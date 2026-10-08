"""Exportação para Excel, seed de demonstração, health e registro de ação do Backoffice.

Saiu do main.py na refatoração de 2026-10-08, sem mudança de comportamento.
"""
from fastapi import APIRouter, Cookie, HTTPException
from pydantic import BaseModel

from app.core.auth import get_session, hash_senha
from app.core.db import get_db

# --- corpo ---
router = APIRouter()


class AcaoBackoffice(BaseModel):
    comando: str
    cliente: str = "Geral"


@router.post("/api/registrar-acao")
def registrar_acao(acao: AcaoBackoffice, faiston_token: str = Cookie(None)):
    return {"sucesso": True, "mensagem": "Ação registrada"}

@router.get("/api/health")
def health(): return {"status": "ok"}

@router.get("/api/exportar")
def exportar_excel(cliente: str = "", data_inicio: str = "", data_fim: str = "", faiston_token: str = Cookie(None)):
    from fastapi.responses import StreamingResponse
    import io
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401, detail="Não autenticado")
    if sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403, detail="Acesso negado")
    conn = get_db()
    if not conn: raise HTTPException(status_code=500, detail="Banco offline")
    try:
        cur = conn.cursor()
        conditions = []
        params = []
        if cliente:
            conditions.append("t.cliente = %s")
            params.append(cliente)
        if data_inicio:
            conditions.append("t.criado_em >= %s")
            params.append(data_inicio + " 00:00:00")
        if data_fim:
            conditions.append("t.criado_em <= %s")
            params.append(data_fim + " 23:59:59")
        filtro = ("WHERE " + " AND ".join(conditions)) if conditions else ""
        cur.execute(f"""SELECT u.nome, t.descricao, t.cliente, t.prioridade, t.status,
            t.segundos, t.criado_em, t.atualizado_em, COALESCE(t.peso,0),
            t.prazo_status, t.justificativa_atraso
            FROM tarefas t JOIN usuarios u ON t.usuario_id = u.id
            {filtro} ORDER BY t.criado_em DESC""", tuple(params))
        rows = cur.fetchall()
        cur.close(); conn.close()

        # Gerar XLSX com openpyxl
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tarefas"

        # Estilos
        header_fill = PatternFill("solid", fgColor="4A00E0")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        alt_fill = PatternFill("solid", fgColor="F8FAFC")
        border = Border(bottom=Side(style='thin', color='E2E8F0'))
        center = Alignment(horizontal='center', vertical='center')

        # Cabeçalho -- data e hora de criação/atualização em colunas separadas
        # (em vez de um texto "AAAA-MM-DD HH:MM" só, mais fácil de ler/filtrar
        # fora do Excel), peso da atividade, e o desfecho do prazo (carimbado
        # só na conclusão) com o motivo do atraso quando ficou fora do prazo.
        headers = ["Funcionário", "Tarefa", "Cliente", "Prioridade", "Peso", "Status", "Horas", "Minutos", "Total (h)",
                   "Criado em", "Hora criação", "Atualizado em", "Hora atualização", "Prazo", "Motivo do atraso"]
        col_widths = [25, 40, 20, 12, 8, 15, 8, 8, 10, 14, 12, 14, 12, 16, 40]
        for col, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 30

        # Dados
        status_map = {"concluido": "Concluído", "em_andamento": "Em Andamento", "aberto": "Aberto"}
        prio_colors = {"Alta": "FFE4E6", "Media": "FEF3C7", "Baixa": "D1FAE5"}
        status_colors = {"concluido": "D1FAE5", "em_andamento": "CFFAFE", "aberto": "F1F5F9"}
        prazo_map = {"dentro": "Dentro do prazo", "fora": "Fora do prazo", "sem_prazo": "Sem prazo definido"}

        DATE_COLS = (10, 12)  # Criado em / Atualizado em
        TIME_COLS = (11, 13)  # Hora criação / Hora atualização
        for i, r in enumerate(rows, 2):
            h = r[5] // 3600
            m = (r[5] % 3600) // 60
            total_h = round(r[5] / 3600, 2)
            status_label = status_map.get(r[4], r[4])
            prazo_label = prazo_map.get(r[9], "")
            row_data = [r[0], r[1], r[2], r[3], r[8], status_label, h, m, total_h,
                r[6].date() if r[6] else None, r[6].time() if r[6] else None,
                r[7].date() if r[7] else None, r[7].time() if r[7] else None,
                prazo_label, r[10] or ""]
            fill = PatternFill("solid", fgColor="FFFFFF") if i % 2 == 0 else alt_fill
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=i, column=col, value=val)
                cell.border = border
                cell.alignment = Alignment(vertical='center')
                if col in DATE_COLS:
                    cell.number_format = "DD/MM/YYYY"
                elif col in TIME_COLS:
                    cell.number_format = "HH:MM"
                # Cor por prioridade, status e prazo
                if col == 4 and r[3] in prio_colors:
                    cell.fill = PatternFill("solid", fgColor=prio_colors[r[3]])
                elif col == 6 and r[4] in status_colors:
                    cell.fill = PatternFill("solid", fgColor=status_colors[r[4]])
                elif col == 14 and r[9] == "fora":
                    cell.fill = PatternFill("solid", fgColor="FFE4E6")
                    cell.font = Font(color="C02234", bold=True)
                else:
                    cell.fill = fill
            ws.row_dimensions[i].height = 22

        # Totais
        total_row = len(rows) + 2
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=7, value=sum(r[5]//3600 for r in rows)).font = Font(bold=True)
        ws.cell(row=total_row, column=9, value=round(sum(r[5] for r in rows)/3600, 2)).font = Font(bold=True)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"faiston_tarefas{'_'+cliente if cliente else ''}{'_'+data_inicio if data_inicio else ''}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.delete("/api/limpar-seed")
def limpar_seed(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] != "admin": raise HTTPException(status_code=403, detail="Apenas admin")
    conn = get_db()
    if not conn: raise HTTPException(status_code=500, detail="Banco offline")
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM tarefas WHERE usuario_id IN (SELECT id FROM usuarios WHERE usuario IN ('mariana','joao','carlos','fernanda','thiago'))")
        tarefas = cur.rowcount
        cur.execute("DELETE FROM tarefas WHERE descricao LIKE '%[TESTE]%'")
        tarefas += cur.rowcount
        cur.execute("DELETE FROM usuarios WHERE usuario IN ('mariana','joao','carlos','fernanda','thiago')")
        usuarios = cur.rowcount
        # Remove também projetos, clientes e histórico de teste (namespaced [TESTE])
        cur.execute("DELETE FROM tarefa_historico WHERE COALESCE(tarefa_desc,'') LIKE '%[TESTE]%' OR COALESCE(projeto_nome,'') LIKE '%[TESTE]%'")
        cur.execute("DELETE FROM projetos WHERE nome LIKE '%[TESTE]%'")
        cur.execute("DELETE FROM clientes WHERE nome LIKE '%[TESTE]%'")
        conn.commit()
        return {"ok": True, "tarefas_removidas": tarefas, "usuarios_removidos": usuarios}
    except Exception as e:
        conn.rollback()
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/seed-dados")
def seed_dados(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] != "admin": raise HTTPException(status_code=403, detail="Apenas admin")
    import random
    from datetime import datetime, timedelta
    conn = get_db()
    if not conn: raise HTTPException(status_code=500, detail="Banco offline")
    try:
        cur = conn.cursor()
        # Garante colunas usadas pelo seed (migração lazy normalmente roda no GET /api/tarefas)
        cur.execute("ALTER TABLE tarefas ADD COLUMN IF NOT EXISTS projeto_id INTEGER REFERENCES projetos(id)")
        cur.execute("ALTER TABLE tarefas ADD COLUMN IF NOT EXISTS data_prazo DATE")
        cur.execute("ALTER TABLE tarefas ADD COLUMN IF NOT EXISTS hora_prazo TIME")
        # Limpa qualquer dado de teste anterior (ordem respeita as FKs)
        cur.execute("DELETE FROM tarefas WHERE descricao LIKE '%[TESTE]%'")
        cur.execute("DELETE FROM usuarios WHERE usuario IN ('mariana','joao','carlos','fernanda','thiago')")
        cur.execute("DELETE FROM tarefa_historico WHERE COALESCE(tarefa_desc,'') LIKE '%[TESTE]%' OR COALESCE(projeto_nome,'') LIKE '%[TESTE]%'")
        cur.execute("DELETE FROM projetos WHERE nome LIKE '%[TESTE]%'")
        cur.execute("DELETE FROM clientes WHERE nome LIKE '%[TESTE]%'")

        funcionarios = [
            ("mariana","faiston123","Mariana Silva","funcionario"),
            ("joao","faiston123","João Henrique","funcionario"),
            ("carlos","faiston123","Carlos Eduardo","funcionario"),
            ("fernanda","faiston123","Fernanda Lima","funcionario"),
            ("thiago","faiston123","Thiago Rocha","funcionario"),
        ]
        ids = {}
        for usuario, senha, nome, perfil in funcionarios:
            cur.execute("""INSERT INTO usuarios (usuario, senha_hash, nome, perfil, primeiro_acesso)
                VALUES (%s,%s,%s,%s,FALSE) ON CONFLICT (usuario) DO UPDATE SET nome=%s RETURNING id""",
                (usuario, hash_senha(senha), nome, perfil, nome))
            ids[nome] = cur.fetchone()[0]

        clientes = ["NTT","Arcos Dourados","Zamp","Telcoweb","VIVO VITA"]
        prioridades_peso = ["Critica","Alta","Alta","Media","Media","Media","Baixa"]

        # Clientes e projetos de teste (namespaced [TESTE] → limpeza segura, não toca dado real)
        clientes_proj = {
            "[TESTE] NTT":            ["[TESTE] NOC 24x7", "[TESTE] Backbone MPLS"],
            "[TESTE] Zamp":           ["[TESTE] Rollout de Lojas"],
            "[TESTE] Arcos Dourados": ["[TESTE] Field Services"],
        }
        projetos_ref = []  # (projeto_id, projeto_nome, cliente_nome)
        for cli_nome, projs in clientes_proj.items():
            cur.execute("INSERT INTO clientes (nome) VALUES (%s) ON CONFLICT (nome) DO UPDATE SET nome=EXCLUDED.nome RETURNING id", (cli_nome,))
            cli_id = cur.fetchone()[0]
            for pnome in projs:
                cur.execute("INSERT INTO projetos (cliente_id, nome) VALUES (%s,%s) RETURNING id", (cli_id, pnome))
                projetos_ref.append((cur.fetchone()[0], pnome, cli_nome))
        descricoes = [
            "[TESTE] Abertura de chamado no NOC",
            "[TESTE] Acompanhamento de incidente crítico",
            "[TESTE] Configuração de switch core",
            "[TESTE] Monitoramento de links MPLS",
            "[TESTE] Troca de equipamento defeituoso",
            "[TESTE] Atualização de firmware",
            "[TESTE] Relatório de disponibilidade mensal",
            "[TESTE] Escalada para fornecedor",
            "[TESTE] Revisão de topologia de rede",
            "[TESTE] Acionamento de parceiro técnico",
            "[TESTE] Documentação de circuito",
            "[TESTE] Teste de failover",
            "[TESTE] Análise de log de erros",
            "[TESTE] Validação de SLA",
            "[TESTE] Suporte remoto ao cliente",
            "[TESTE] Instalação de CPE",
            "[TESTE] Diagnóstico de latência",
            "[TESTE] Follow-up de chamado crítico",
        ]
        now = datetime.now()
        total = 0
        criadas = []  # (tid, desc, cliente, projeto_id, projeto_nome, owner_nome)

        def add_tarefa(owner_nome, desc, status, segundos, prioridade, dias_atras, com_projeto=True):
            # ~75% das tarefas vão para um projeto de teste (pra exercitar o agrupamento
            # por projeto); o restante fica "Sem projeto" (cliente avulso).
            if com_projeto and projetos_ref and random.random() < 0.75:
                pid, pnome, cli = random.choice(projetos_ref)
            else:
                pid, pnome, cli = None, None, random.choice(clientes)
            ts = now - timedelta(days=dias_atras)
            cur.execute("""INSERT INTO tarefas (usuario_id,descricao,cliente,prioridade,status,segundos,projeto_id,criado_em,atualizado_em)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id""",
                        (ids[owner_nome], desc, cli, prioridade, status, segundos, pid, ts, now))
            tid = cur.fetchone()[0]
            criadas.append((tid, desc, cli, pid, pnome, owner_nome))
            return tid

        # Mariana — sobrecarregada, muitos abertos antigos (IA deve alertar)
        for i in range(12):
            add_tarefa("Mariana Silva", random.choice(descricoes), "aberto", 0,
                       random.choice(["Alta","Critica"]), random.randint(8, 20))
            total += 1
        # João — lento, tickets em andamento há muito tempo
        for i in range(8):
            add_tarefa("João Henrique", random.choice(descricoes),
                       random.choice(["em_andamento","em_andamento","aberto"]),
                       random.randint(600, 3600), random.choice(prioridades_peso), random.randint(5, 15))
            total += 1
        # Carlos — equilibrado, maioria concluído
        for i in range(10):
            status = random.choice(["concluido","concluido","concluido","em_andamento","aberto"])
            segundos = random.randint(1800, 7200) if status == "concluido" else random.randint(600, 3600)
            add_tarefa("Carlos Eduardo", random.choice(descricoes), status, segundos,
                       random.choice(prioridades_peso), random.randint(1, 7))
            total += 1
        # Fernanda — poucos tickets, bem resolvidos
        for i in range(5):
            add_tarefa("Fernanda Lima", random.choice(descricoes),
                       random.choice(["concluido","concluido","aberto"]), random.randint(1800, 5400),
                       "Media", random.randint(1, 5))
            total += 1
        # Thiago — vários críticos abertos
        for i in range(7):
            add_tarefa("Thiago Rocha", random.choice(descricoes),
                       random.choice(["aberto","aberto","em_andamento"]), 0,
                       random.choice(["Critica","Alta"]), random.randint(2, 10))
            total += 1

        # ── Histórico de "hoje" (NOW()) para popular o resumo diário ──
        def add_hist(tid, desc, pid, pnome, autor, acao, campo, antigo, novo):
            cur.execute("""INSERT INTO tarefa_historico
                (tarefa_id, tarefa_desc, autor_id, autor_nome, acao, campo, valor_antigo, valor_novo,
                 projeto_id, projeto_nome, time_tarefa, criado_em)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'Projetos', NOW())""",
                (tid, desc, ids.get(autor), autor, acao, campo, antigo, novo, pid, pnome))

        com_proj = [c for c in criadas if c[3]]
        random.shuffle(com_proj)
        hist = 0
        for c in com_proj[:6]:
            add_hist(c[0], c[1], c[3], c[4], c[5], "criou", "tarefa", None, c[1]); hist += 1
        for c in com_proj[6:10]:
            add_hist(c[0], c[1], c[3], c[4], c[5], "editou", "status", "aberto", "em_andamento"); hist += 1
            add_hist(c[0], c[1], c[3], c[4], c[5], "editou", "prioridade", "Media", "Alta"); hist += 1
        if com_proj:
            c = com_proj[0]
            add_hist(c[0], c[1], c[3], c[4], c[5], "editou", "data_prazo", None,
                     (now + timedelta(days=1)).strftime("%Y-%m-%d")); hist += 1
        if len(com_proj) > 10:
            c = com_proj[10]
            add_hist(c[0], c[1], c[3], c[4], c[5], "excluiu", "tarefa", c[1], None); hist += 1

        conn.commit(); cur.close(); conn.close()
        return {
            "sucesso": True,
            "tarefas_criadas": total,
            "projetos_criados": len(projetos_ref),
            "historico_criado": hist,
            "usuarios": {u[0]: "senha: faiston123" for u in funcionarios}
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
