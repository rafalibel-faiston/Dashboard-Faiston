"""Forecast / P&L (adaptação da planilha FORECAST2026).

Saiu do main.py na refatoração de 2026-10-08, sem mudança de comportamento.
"""
from datetime import date, datetime
from typing import Optional
import io

from fastapi import APIRouter, Cookie, File, HTTPException, UploadFile
from pydantic import BaseModel

from app.core.auth import get_session
from app.core.db import get_db


router = APIRouter()


# ─────────────────────────────────────────────
#  FORECAST / P&L  (adaptação da planilha FORECAST2026)
#  RESUMO  = consolidado de margem por projeto
#  MENSAL  = P&L por período (custos detalhados)
# ─────────────────────────────────────────────

# Ligado por _garantir_tabelas_na_subida (main.py) depois do commit -- mesma
# lógica de _FINANCEIRO_TABELAS_OK em app/financeiro.
_FORECAST_TABELAS_OK = False


def _ensure_forecast_tables(cur):
    if _FORECAST_TABELAS_OK:
        return
    cur.execute("""
        CREATE TABLE IF NOT EXISTS forecast_projetos (
            id SERIAL PRIMARY KEY,
            codigo VARCHAR(40) NOT NULL,
            cliente VARCHAR(150) DEFAULT '',
            cliente_final VARCHAR(150) DEFAULT '',
            projeto VARCHAR(300) DEFAULT '',
            data_inicio DATE,
            data_fim DATE,
            status VARCHAR(40) DEFAULT '',
            acum_fat NUMERIC(16,2) DEFAULT 0,
            pct_fat NUMERIC(10,6) DEFAULT 0,
            pend_fat NUMERIC(16,2) DEFAULT 0,
            total NUMERIC(16,2) DEFAULT 0,
            gross_revenue NUMERIC(16,2) DEFAULT 0,
            tax NUMERIC(10,6) DEFAULT 0,
            imp NUMERIC(16,2) DEFAULT 0,
            net_revenue NUMERIC(16,2) DEFAULT 0,
            acum_custo NUMERIC(16,2) DEFAULT 0,
            valor_dm NUMERIC(16,2) DEFAULT 0,
            pct_margem NUMERIC(10,6) DEFAULT 0,
            pend_custo NUMERIC(16,2) DEFAULT 0,
            obs TEXT DEFAULT '',
            atualizado_em TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS forecast_mensal (
            id SERIAL PRIMARY KEY,
            periodo VARCHAR(20) NOT NULL,
            periodo_key VARCHAR(20) NOT NULL,
            codigo VARCHAR(40) NOT NULL,
            cliente VARCHAR(150) DEFAULT '',
            cliente_final VARCHAR(150) DEFAULT '',
            projeto VARCHAR(300) DEFAULT '',
            data_inicio DATE,
            status VARCHAR(40) DEFAULT '',
            gross_revenue NUMERIC(16,2) DEFAULT 0,
            tax NUMERIC(10,6) DEFAULT 0,
            imp NUMERIC(16,2) DEFAULT 0,
            net_revenue NUMERIC(16,2) DEFAULT 0,
            custo_folha NUMERIC(16,2) DEFAULT 0,
            logistica NUMERIC(16,2) DEFAULT 0,
            material NUMERIC(16,2) DEFAULT 0,
            vlr_parceiros NUMERIC(16,2) DEFAULT 0,
            spare_parts NUMERIC(16,2) DEFAULT 0,
            outros NUMERIC(16,2) DEFAULT 0,
            total_custo NUMERIC(16,2) DEFAULT 0,
            atualizado_em TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_forecast_mensal_periodo ON forecast_mensal(periodo_key)")


_FC_MESES = {'JAN':1,'FEV':2,'MAR':3,'ABR':4,'MAI':5,'JUN':6,'JUL':7,'AGO':8,'SET':9,'OUT':10,'NOV':11,'DEZ':12}

_FC_RESUMO_MAP = {
    'ID':'codigo','CLIENTE':'cliente','CLIENTE FINAL':'cliente_final','PROJETO':'projeto',
    'INICIO':'data_inicio','FIM':'data_fim','STATUS':'status','ACUM FAT':'acum_fat','% FAT':'pct_fat',
    'PEND FAT':'pend_fat','TOTAL':'total','GROSS REVENUE':'gross_revenue','TAX':'tax','IMP':'imp',
    'NET REVENUE':'net_revenue','ACUM CUSTO':'acum_custo','VALOR DM':'valor_dm','%':'pct_margem',
    'PEND CUSTO':'pend_custo','OBS':'obs',
}
_FC_MENSAL_MAP = {
    'ID':'codigo','CLIENTE':'cliente','CLIENTE FINAL':'cliente_final','PROJETO':'projeto',
    'INICIO':'data_inicio','STATUS':'status','GROSS REVENUE':'gross_revenue','TAX':'tax','IMP':'imp',
    'NET REVENUE':'net_revenue','CUSTO FOLHA':'custo_folha','LOGISTICA':'logistica','MATERIAL':'material',
    'VLR PARCEIROS':'vlr_parceiros','SPARE PARTS':'spare_parts','OUTROS':'outros','TOTAL CUSTO':'total_custo',
}
_FC_NUM_FIELDS = {'acum_fat','pct_fat','pend_fat','total','gross_revenue','tax','imp','net_revenue',
                  'acum_custo','valor_dm','pct_margem','pend_custo','custo_folha','logistica','material',
                  'vlr_parceiros','spare_parts','outros','total_custo'}
_FC_DATE_FIELDS = {'data_inicio','data_fim'}

def _fc_norm(h):
    import re as _re
    if h is None: return ""
    s = str(h).strip().upper().replace('.', ' ')
    return _re.sub(r'\s+', ' ', s).strip()

def _fc_num(v):
    if v is None or v == '': return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace('R$', '').replace('.', '').replace(',', '.').strip())
    except Exception: return 0.0

def _fc_dt(v):
    if isinstance(v, (datetime, date)): return v.isoformat()[:10]
    return None

def _fc_find_header(rows):
    for i, r in enumerate(rows[:15]):
        if r and _fc_norm(r[0]) == 'ID':
            return i
    return None

def _parse_forecast_workbook(content):
    """Lê uma planilha no formato FORECAST (aba RESUMO + abas mensais/anuais)."""
    import re as _re
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    resumo, mensal, periodos = [], [], []
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        hi = _fc_find_header(rows)
        if hi is None: continue
        up = sname.strip().upper()
        m = _re.match(r'^([A-Z]{3})(\d{2})$', up)
        if up == 'RESUMO':
            kind, mapping, periodo, pkey = 'resumo', _FC_RESUMO_MAP, None, None
        elif m and m.group(1) in _FC_MESES:
            kind, mapping = 'mensal', _FC_MENSAL_MAP
            periodo = sname.strip(); pkey = f"20{m.group(2)}-{_FC_MESES[m.group(1)]:02d}"
        elif _re.match(r'^\d{4}$', up):
            kind, mapping = 'mensal', _FC_MENSAL_MAP
            periodo = sname.strip(); pkey = up
        else:
            continue
        headers = rows[hi]
        col = {}
        for idx, h in enumerate(headers):
            key = mapping.get(_fc_norm(h))
            if key and key not in col: col[key] = idx
        if 'codigo' not in col: continue
        cnt = 0
        for r in rows[hi + 1:]:
            codigo = r[col['codigo']] if col['codigo'] < len(r) else None
            if codigo is None or not str(codigo).strip(): continue
            rec = {'periodo': periodo, 'periodo_key': pkey}
            for field, idx in col.items():
                v = r[idx] if idx < len(r) else None
                if field in _FC_NUM_FIELDS: rec[field] = _fc_num(v)
                elif field in _FC_DATE_FIELDS: rec[field] = _fc_dt(v)
                else: rec[field] = (str(v).strip() if v is not None else '')
            (resumo if kind == 'resumo' else mensal).append(rec)
            cnt += 1
        if kind == 'mensal': periodos.append({'periodo_key': pkey, 'periodo': periodo, 'linhas': cnt})
    return resumo, mensal, periodos

@router.post("/api/forecast/importar")
async def forecast_importar(file: UploadFile = File(...), faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox; openpyxl = _ox; _OPENPYXL_OK = True
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        content = await file.read()
        resumo, mensal, periodos = _parse_forecast_workbook(content)
        if not resumo and not mensal:
            raise HTTPException(status_code=400, detail="Nenhuma aba reconhecida (esperado RESUMO + abas mensais como JAN26, 2025).")
        cur = conn.cursor()
        _ensure_forecast_tables(cur)
        # Substitui toda a base do forecast (a planilha é a fonte da verdade)
        cur.execute("TRUNCATE forecast_projetos, forecast_mensal RESTART IDENTITY")
        rcols = ['codigo','cliente','cliente_final','projeto','data_inicio','data_fim','status',
                 'acum_fat','pct_fat','pend_fat','total','gross_revenue','tax','imp','net_revenue',
                 'acum_custo','valor_dm','pct_margem','pend_custo','obs']
        for r in resumo:
            cur.execute(
                f"INSERT INTO forecast_projetos ({','.join(rcols)}) VALUES ({','.join(['%s']*len(rcols))})",
                tuple(r.get(c) for c in rcols))
        mcols = ['periodo','periodo_key','codigo','cliente','cliente_final','projeto','data_inicio','status',
                 'gross_revenue','tax','imp','net_revenue','custo_folha','logistica','material',
                 'vlr_parceiros','spare_parts','outros','total_custo']
        for r in mensal:
            cur.execute(
                f"INSERT INTO forecast_mensal ({','.join(mcols)}) VALUES ({','.join(['%s']*len(mcols))})",
                tuple(r.get(c) for c in mcols))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "resumo": len(resumo), "mensal": len(mensal),
                "periodos": sorted([p['periodo_key'] for p in periodos])}
    except HTTPException: raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Erro ao importar forecast: {str(e)}")

def _fc_to_float(v): return float(v) if v is not None else 0.0

@router.get("/api/forecast/resumo")
def forecast_get_resumo(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_forecast_tables(cur); conn.commit()
        cur.execute("""
            SELECT id, codigo, cliente, cliente_final, projeto, data_inicio, data_fim, status,
                   acum_fat, pct_fat, pend_fat, total, gross_revenue, tax, imp, net_revenue,
                   acum_custo, valor_dm, pct_margem, pend_custo, obs
            FROM forecast_projetos
            ORDER BY total DESC NULLS LAST, codigo
        """)
        rows = cur.fetchall()
        cur.execute("SELECT MAX(atualizado_em) FROM forecast_projetos")
        atualizado = cur.fetchone()[0]
        cur.close(); conn.close()
        projetos = [{
            "id": r[0], "codigo": r[1], "cliente": r[2], "cliente_final": r[3], "projeto": r[4],
            "data_inicio": str(r[5]) if r[5] else None, "data_fim": str(r[6]) if r[6] else None,
            "status": r[7], "acum_fat": _fc_to_float(r[8]), "pct_fat": _fc_to_float(r[9]),
            "pend_fat": _fc_to_float(r[10]), "total": _fc_to_float(r[11]), "gross_revenue": _fc_to_float(r[12]),
            "tax": _fc_to_float(r[13]), "imp": _fc_to_float(r[14]), "net_revenue": _fc_to_float(r[15]),
            "acum_custo": _fc_to_float(r[16]), "valor_dm": _fc_to_float(r[17]), "pct_margem": _fc_to_float(r[18]),
            "pend_custo": _fc_to_float(r[19]), "obs": r[20] or "",
        } for r in rows]
        tot = {
            "total": sum(p["total"] for p in projetos),
            "gross_revenue": sum(p["gross_revenue"] for p in projetos),
            "net_revenue": sum(p["net_revenue"] for p in projetos),
            "acum_fat": sum(p["acum_fat"] for p in projetos),
            "pend_fat": sum(p["pend_fat"] for p in projetos),
            "acum_custo": sum(p["acum_custo"] for p in projetos),
            "pend_custo": sum(p["pend_custo"] for p in projetos),
            "valor_dm": sum(p["valor_dm"] for p in projetos),
        }
        tot["margem_pct"] = (tot["valor_dm"] / tot["net_revenue"]) if tot["net_revenue"] else 0.0
        return {"projetos": projetos, "totais": tot,
                "atualizado_em": str(atualizado)[:19] if atualizado else None}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/forecast/periodos")
def forecast_periodos(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_forecast_tables(cur); conn.commit()
        cur.execute("""
            SELECT periodo_key, MAX(periodo), COUNT(*),
                   SUM(net_revenue), SUM(total_custo)
            FROM forecast_mensal GROUP BY periodo_key ORDER BY periodo_key
        """)
        rows = cur.fetchall()
        cur.close(); conn.close()
        return [{"periodo_key": r[0], "periodo": r[1], "linhas": r[2],
                 "net_revenue": _fc_to_float(r[3]), "total_custo": _fc_to_float(r[4]),
                 "margem": _fc_to_float(r[3]) - _fc_to_float(r[4])} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/forecast/mensal")
def forecast_mensal(periodo: str = "", faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_forecast_tables(cur); conn.commit()
        if not periodo:
            cur.execute("SELECT periodo_key FROM forecast_mensal ORDER BY periodo_key DESC LIMIT 1")
            row = cur.fetchone()
            periodo = row[0] if row else ""
        cur.execute("""
            SELECT codigo, cliente, cliente_final, projeto, status,
                   gross_revenue, tax, imp, net_revenue,
                   custo_folha, logistica, material, vlr_parceiros, spare_parts, outros, total_custo
            FROM forecast_mensal WHERE periodo_key=%s
            ORDER BY net_revenue DESC NULLS LAST, codigo
        """, (periodo,))
        rows = cur.fetchall()
        cur.close(); conn.close()
        linhas = [{
            "codigo": r[0], "cliente": r[1], "cliente_final": r[2], "projeto": r[3], "status": r[4],
            "gross_revenue": _fc_to_float(r[5]), "tax": _fc_to_float(r[6]), "imp": _fc_to_float(r[7]),
            "net_revenue": _fc_to_float(r[8]), "custo_folha": _fc_to_float(r[9]), "logistica": _fc_to_float(r[10]),
            "material": _fc_to_float(r[11]), "vlr_parceiros": _fc_to_float(r[12]), "spare_parts": _fc_to_float(r[13]),
            "outros": _fc_to_float(r[14]), "total_custo": _fc_to_float(r[15]),
            "margem": _fc_to_float(r[8]) - _fc_to_float(r[15]),
        } for r in rows]
        campos_custo = ["custo_folha","logistica","material","vlr_parceiros","spare_parts","outros","total_custo"]
        tot = {k: sum(l[k] for l in linhas) for k in
               ["gross_revenue","imp","net_revenue"] + campos_custo}
        tot["margem"] = tot["net_revenue"] - tot["total_custo"]
        tot["margem_pct"] = (tot["margem"] / tot["net_revenue"]) if tot["net_revenue"] else 0.0
        return {"periodo": periodo, "linhas": linhas, "totais": tot}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))


class ForecastProjetoModel(BaseModel):
    codigo: str
    cliente: str = ""
    cliente_final: str = ""
    projeto: str = ""
    data_inicio: Optional[str] = None
    data_fim: Optional[str] = None
    status: str = "ANDAMENTO"
    total: float = 0.0          # valor de contrato (= gross revenue)
    tax: float = 0.1125         # alíquota de imposto (fração)
    acum_fat: float = 0.0       # faturamento acumulado
    acum_custo: float = 0.0     # custo acumulado
    pend_custo: float = 0.0     # custo pendente (previsto)
    obs: str = ""

def _fc_calcular(m: "ForecastProjetoModel"):
    """Deriva impostos, net revenue e margem a partir dos campos informados,
    seguindo a mesma lógica da planilha FORECAST."""
    total = m.total or 0.0
    gross = total
    tax = m.tax or 0.0
    imp = gross * tax
    net = gross - imp
    acum_custo = m.acum_custo or 0.0
    valor_dm = net - acum_custo
    pct_margem = (valor_dm / net) if net else 0.0
    acum_fat = m.acum_fat or 0.0
    pct_fat = (acum_fat / total) if total else 0.0
    pend_fat = total - acum_fat
    return {
        "gross_revenue": gross, "imp": imp, "net_revenue": net,
        "valor_dm": valor_dm, "pct_margem": pct_margem,
        "pct_fat": pct_fat, "pend_fat": pend_fat,
    }

@router.post("/api/forecast/projeto")
def forecast_criar_projeto(m: ForecastProjetoModel, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    if not (m.codigo or "").strip(): raise HTTPException(status_code=400, detail="Informe o ID do projeto.")
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_forecast_tables(cur)
        d = _fc_calcular(m)
        cur.execute("""
            INSERT INTO forecast_projetos
                (codigo, cliente, cliente_final, projeto, data_inicio, data_fim, status,
                 acum_fat, pct_fat, pend_fat, total, gross_revenue, tax, imp, net_revenue,
                 acum_custo, valor_dm, pct_margem, pend_custo, obs, atualizado_em)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW())
            RETURNING id
        """, (m.codigo.strip(), m.cliente, m.cliente_final, m.projeto, m.data_inicio or None,
              m.data_fim or None, m.status, m.acum_fat, d["pct_fat"], d["pend_fat"], m.total,
              d["gross_revenue"], m.tax, d["imp"], d["net_revenue"], m.acum_custo, d["valor_dm"],
              d["pct_margem"], m.pend_custo, m.obs))
        new_id = cur.fetchone()[0]
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "id": new_id}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.put("/api/forecast/projeto/{fid}")
def forecast_atualizar_projeto(fid: int, m: ForecastProjetoModel, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    if not (m.codigo or "").strip(): raise HTTPException(status_code=400, detail="Informe o ID do projeto.")
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_forecast_tables(cur)
        d = _fc_calcular(m)
        cur.execute("""
            UPDATE forecast_projetos SET
                codigo=%s, cliente=%s, cliente_final=%s, projeto=%s, data_inicio=%s, data_fim=%s,
                status=%s, acum_fat=%s, pct_fat=%s, pend_fat=%s, total=%s, gross_revenue=%s, tax=%s,
                imp=%s, net_revenue=%s, acum_custo=%s, valor_dm=%s, pct_margem=%s, pend_custo=%s,
                obs=%s, atualizado_em=NOW()
            WHERE id=%s
        """, (m.codigo.strip(), m.cliente, m.cliente_final, m.projeto, m.data_inicio or None,
              m.data_fim or None, m.status, m.acum_fat, d["pct_fat"], d["pend_fat"], m.total,
              d["gross_revenue"], m.tax, d["imp"], d["net_revenue"], m.acum_custo, d["valor_dm"],
              d["pct_margem"], m.pend_custo, m.obs, fid))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.delete("/api/forecast/projeto/{fid}")
def forecast_deletar_projeto(fid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo", "diretor"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM forecast_projetos WHERE id=%s", (fid,))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))
