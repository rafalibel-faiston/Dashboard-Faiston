"""Importação de planilha de cronograma/atividades para o Status de Campo.

Saiu do main.py na refatoração de 2026-10-08, sem mudança de comportamento.
"""
from datetime import date, datetime
from typing import Optional
import io

from fastapi import APIRouter, Cookie, File, Form, HTTPException, UploadFile

from app.core.acesso import _eh_backoffice, _pode_ver_status_report
from app.core.auth import get_session
from app.core.db import get_db

# --- corpo ---
router = APIRouter()


# ── Importação de planilha de cronograma/atividades ─────────────────────────
# Mapeia só as colunas que já existem no sistema (cliente, subprojeto,
# cidade/UF, técnico, agendamento/horário, status, N2 responsável, ticket,
# observação) -- ignora de propósito colunas sensíveis/sem equivalente
# (RG/CPF/telefone, valores financeiros, KM, HE, seriais, estoque etc.).
_SC_IMPORT_HEADERS = {
    'CLIENTE': 'cliente', 'PROJETO': 'projeto', 'SUBPROJETO': 'subprojeto',
    'RESPONSAVEL': 'n2_responsavel', 'LOCALIDADE/ITASK': 'site', 'LOCALIDADE': 'site',
    'ENDERECO': 'endereco', 'CIDADE': 'cidade', 'UF': 'uf',
    'AGENDAMENTO': 'data', 'HORARIO': 'horario_agendado', 'TECNICO': 'tecnico',
    'TICKET ATENDIMENTO': 'ticket', 'STATUS ATIVIDADE': 'status_raw', 'OBSERVACAO': 'observacoes',
    # Dialeto "cronograma de parceiro" (ex.: VITA/Arcos Dourados) -- é a
    # planilha que o time já usa no dia a dia, de um único cliente, sem
    # coluna CLIENTE (o cliente vem do seletor no modal de importar, ver
    # `cliente_id`) (2026-08-18).
    'NOME': 'site', 'SIGLA': 'site_sigla', 'DATA': 'data', 'STATUS': 'status_raw',
    'TICKET': 'ticket', 'ATIVIDADE': 'subprojeto',
}

def _sc_strip_acentos(s):
    import unicodedata
    return ''.join(c for c in unicodedata.normalize('NFKD', str(s)) if not unicodedata.combining(c))

def _sc_norm_header(h):
    import re as _re
    if h is None: return ""
    s = _sc_strip_acentos(str(h)).strip().upper()
    return _re.sub(r'\s+', ' ', s)

def _sc_norm_nome(s):
    import re as _re
    s = _sc_strip_acentos(str(s or '')).strip().upper()
    s = _re.sub(r'[^A-Z0-9 ]', ' ', s)
    return _re.sub(r'\s+', ' ', s).strip()

def _sc_mapear_status(raw):
    u = _sc_strip_acentos(str(raw or '')).upper()
    if 'CONCLUID' in u: return 'concluido'
    if 'PARCIAL' in u: return 'parcial'
    if 'IMPRODUTIV' in u: return 'improdutiva_cliente'
    if 'CANCELAD' in u: return 'cancelado'
    if 'ANDAMENTO' in u: return 'em_andamento'
    if 'AGENDAD' in u: return 'agendado'
    # "REAGENDAR" (sem sufixo -O/-A) e "AGUARD. AGENDAMENTO" não batem com
    # 'AGENDAD' acima -- planilhas como a da VITA/Arcos Dourados usam essas
    # variações pra dizer a mesma coisa: atividade pendente de agendamento.
    if 'REAGENDAR' in u: return 'agendado'
    if 'AGUARD' in u: return 'agendado'
    return None

def _sc_find_sheet_and_header(wb):
    """Escaneia todas as abas procurando a que tem mais colunas reconhecidas
    (CLIENTE, STATUS ATIVIDADE, TECNICO etc.) nas primeiras linhas -- a
    planilha real pode ter várias abas de resumo/tabela dinâmica junto,
    só a aba com o log linha-a-linha interessa."""
    melhor = None
    for sname in wb.sheetnames:
        ws = wb[sname]
        for i, r in enumerate(ws.iter_rows(max_row=10, values_only=True)):
            normed = {_sc_norm_header(c) for c in r if c is not None}
            hits = len(normed & set(_SC_IMPORT_HEADERS.keys()))
            if hits >= 3 and (melhor is None or hits > melhor[2]):
                melhor = (sname, i, hits)
    return melhor

@router.post("/api/status-campo/importar-planilha")
async def importar_planilha_status_campo(file: UploadFile = File(...), cliente_id: Optional[int] = Form(None),
                                          somente_pendentes: bool = Form(True),
                                          faiston_token: str = Cookie(None)):
    """Importa atividades de uma planilha externa (ex.: cronograma geral),
    mapeando só as colunas que já existem no sistema. Por padrão, Cliente/
    Projeto da planilha são casados por nome aproximado contra os clientes
    já cadastrados -- o que não bate fica de fora e é reportado, não cria
    cliente novo sozinho nem adivinha. Planilhas de um cliente só (ex.:
    cronograma de parceiro tipo VITA/Arcos Dourados, sem coluna CLIENTE)
    passam `cliente_id` explícito no upload -- todas as linhas vão pra esse
    cliente, sem tentar casar nome nenhum (2026-08-18)."""
    sess = get_session(faiston_token)
    if not sess or (sess["perfil"] not in ("admin", "gestor", "demo", "diretor") and not _eh_backoffice(sess)): raise HTTPException(status_code=403)
    if not _pode_ver_status_report(sess): raise HTTPException(status_code=403, detail="Status Report é restrito ao time de Projetos")
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox; openpyxl = _ox; _OPENPYXL_OK = True
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        achado = _sc_find_sheet_and_header(wb)
        if not achado:
            raise HTTPException(status_code=400,
                detail="Nenhuma aba reconhecida (esperado colunas como CLIENTE, STATUS ATIVIDADE, TECNICO etc.)")
        sname, hi, _ = achado
        rows = list(wb[sname].iter_rows(values_only=True))
        headers = rows[hi]
        col = {}
        for idx, h in enumerate(headers):
            key = _SC_IMPORT_HEADERS.get(_sc_norm_header(h))
            if key and key not in col: col[key] = idx

        cur = conn.cursor()
        cur.execute("SELECT id, nome FROM clientes WHERE ativo = TRUE")
        clientes_norm = [(_sc_norm_nome(nome), cid) for cid, nome in cur.fetchall()]

        cliente_fixo = None
        if cliente_id is not None:
            if cliente_id not in {cid for _, cid in clientes_norm}:
                raise HTTPException(status_code=400, detail="Cliente informado não encontrado ou inativo")
            cliente_fixo = cliente_id

        def buscar_cliente(cliente_raw, projeto_raw):
            if cliente_fixo is not None:
                return cliente_fixo
            for cand in (cliente_raw, projeto_raw, f"{cliente_raw} {projeto_raw}".strip()):
                nn = _sc_norm_nome(cand)
                if not nn: continue
                for norm_nome, cid in clientes_norm:
                    if norm_nome and (nn == norm_nome or nn in norm_nome or norm_nome in nn):
                        return cid
            return None

        # Cada linha da planilha só traz o nome do técnico digitado por
        # alguém, igual "Nova atividade" -- sem isso, tecnico_id nunca era
        # preenchido no import, só o texto solto (2026-08-18, a pedido do
        # usuário: "colocar o técnico na atividade de acordo com a base").
        # Nome exato (não substring, ao contrário de cliente) porque nome de
        # pessoa é fácil de dar falso positivo por trecho em comum.
        cur.execute("SELECT id, nome, estado FROM tecnicos WHERE ativo = TRUE")
        tecnicos_norm = [(_sc_norm_nome(nome), (estado or '').strip().upper()[:2], tid)
                          for tid, nome, estado in cur.fetchall()]

        def buscar_tecnico(nome_raw, uf_raw):
            nn = _sc_norm_nome(nome_raw)
            if not nn: return None
            candidatos = [tid for norm_nome, _uf, tid in tecnicos_norm if norm_nome == nn]
            if not candidatos: return None
            uf = (uf_raw or '').strip().upper()[:2]
            if uf:
                na_uf = [tid for norm_nome, _uf, tid in tecnicos_norm if norm_nome == nn and _uf == uf]
                if na_uf: return na_uf[0]
            return candidatos[0]

        def get(r, field):
            idx = col.get(field)
            v = r[idx] if idx is not None and idx < len(r) else None
            return v

        def corta(v, tam):
            return str(v or '').strip()[:tam]

        importadas = 0
        puladas_sem_cliente = {}
        puladas_sem_status = 0
        puladas_sem_data = 0
        puladas_ja_concluidas = 0
        puladas_erro = 0
        tecnicos_nao_encontrados = 0
        for r in rows[hi + 1:]:
            cliente_raw = str(get(r, 'cliente') or '').strip()
            projeto_raw = str(get(r, 'projeto') or '').strip()
            site_raw = str(get(r, 'site') or '').strip()
            # Sem CLIENTE/PROJETO nem SITE a linha está mesmo vazia -- planilhas
            # de cliente único (cliente_fixo) não têm CLIENTE/PROJETO nunca, então
            # SITE (NOME/LOCALIDADE) é o único sinal de linha real que sobra.
            if not cliente_raw and not projeto_raw and not site_raw:
                continue
            data_raw = get(r, 'data')
            if isinstance(data_raw, datetime): data_val = data_raw.date().isoformat()
            elif isinstance(data_raw, date): data_val = data_raw.isoformat()
            else: data_val = None
            if not data_val:
                puladas_sem_data += 1
                continue
            cid = buscar_cliente(cliente_raw, projeto_raw)
            if not cid:
                chave = f"{cliente_raw} {projeto_raw}".strip()
                puladas_sem_cliente[chave] = puladas_sem_cliente.get(chave, 0) + 1
                continue
            status_mapeado = _sc_mapear_status(get(r, 'status_raw'))
            if not status_mapeado:
                puladas_sem_status += 1
                continue
            # Planilha de cronograma real vem com o histórico inteiro junto --
            # sem esse filtro (ligado por padrão), um upload de rotina duplica
            # no banco milhares de atividades já concluídas/canceladas de novo
            # (2026-08-18, a pedido do usuário).
            if somente_pendentes and status_mapeado in ('concluido', 'cancelado'):
                puladas_ja_concluidas += 1
                continue
            horario_raw = get(r, 'horario_agendado')
            horario_val = horario_raw.strftime('%H:%M') if hasattr(horario_raw, 'strftime') else None
            tecnico_raw = get(r, 'tecnico')
            tid = buscar_tecnico(tecnico_raw, get(r, 'uf'))
            if tecnico_raw and not tid:
                tecnicos_nao_encontrados += 1
            # Savepoint por linha -- planilhas reais têm valor fora do
            # padrão de vez em quando (campo longo demais etc.); sem isso,
            # uma linha ruim aborta a transação inteira e nada é salvo.
            cur.execute("SAVEPOINT linha_import")
            try:
                cur.execute("""
                    INSERT INTO status_atividades
                        (cliente_id, data, horario_agendado, tecnico, tecnico_id, n2_responsavel,
                         site_sigla, site_nome, endereco, cidade, uf, subprojeto, ticket, status, observacoes, criado_por)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (cid, data_val, horario_val, corta(tecnico_raw, 150), tid,
                      corta(get(r, 'n2_responsavel'), 150), corta(get(r, 'site_sigla'), 50),
                      corta(get(r, 'site'), 150),
                      str(get(r, 'endereco') or '').strip(), corta(get(r, 'cidade'), 100),
                      corta(get(r, 'uf'), 2), corta(get(r, 'subprojeto'), 150),
                      corta(get(r, 'ticket'), 100), status_mapeado,
                      str(get(r, 'observacoes') or '').strip(), sess["id"]))
                cur.execute("RELEASE SAVEPOINT linha_import")
                importadas += 1
            except Exception:
                cur.execute("ROLLBACK TO SAVEPOINT linha_import")
                puladas_erro += 1
        conn.commit(); cur.close(); conn.close()
        return {
            "sucesso": True, "importadas": importadas, "aba_usada": sname,
            "puladas_erro": puladas_erro,
            "puladas_sem_cliente": [{"nome": k, "ocorrencias": v}
                                     for k, v in sorted(puladas_sem_cliente.items(), key=lambda x: -x[1])],
            "puladas_sem_status": puladas_sem_status, "puladas_sem_data": puladas_sem_data,
            "puladas_ja_concluidas": puladas_ja_concluidas,
            "tecnicos_nao_encontrados": tecnicos_nao_encontrados,
        }
    except HTTPException: raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Erro ao importar planilha: {str(e)}")
