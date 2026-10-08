"""Financeiro: projetos, lançamentos e importação de planilhas.

Saiu do main.py na refatoração de 2026-10-08, sem mudança de comportamento.
"""
from datetime import date
from typing import List, Optional
import csv
import io

from fastapi import APIRouter, Cookie, File, HTTPException, UploadFile
from pydantic import BaseModel

from app.core.acesso import _area_usa_projetos_cur, _exigir_area_com_projeto, _times_validos_cur
from app.core.auth import get_session
from app.core.db import get_db

# --- corpo ---
router = APIRouter()


@router.get("/api/financeiro/resumo")
def financeiro_resumo(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        conn.commit()
        cur.execute("""
            SELECT c.id, c.nome,
                   COALESCE(p_agg.num_projetos, 0) AS num_projetos,
                   COALESCE(p_agg.total_orcamento, 0) AS total_orcamento,
                   COALESCE(l_agg.total_gasto, 0) AS total_gasto
            FROM clientes c
            LEFT JOIN (
                SELECT cliente_id,
                       COUNT(*) AS num_projetos,
                       SUM(orcamento) AS total_orcamento
                FROM projetos
                WHERE ativo = TRUE
                GROUP BY cliente_id
            ) p_agg ON p_agg.cliente_id = c.id
            LEFT JOIN (
                SELECT p.cliente_id, SUM(l.valor) AS total_gasto
                FROM lancamentos l
                JOIN projetos p ON p.id = l.projeto_id AND p.ativo = TRUE
                GROUP BY p.cliente_id
            ) l_agg ON l_agg.cliente_id = c.id
            WHERE c.ativo = TRUE
            ORDER BY COALESCE(l_agg.total_gasto, 0) DESC, c.nome
        """)
        rows = cur.fetchall()
        cur.close(); conn.close()
        return [{"id": r[0], "nome": r[1], "num_projetos": r[2],
                 "orcamento": float(r[3]), "gasto": float(r[4])} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))


# ─────────────────────────────────────────────
#  FINANCEIRO — Projetos e Lançamentos
# ─────────────────────────────────────────────

# Ligado quando a subida do app já garantiu as tabelas (ver
# _garantir_tabelas_na_subida). A partir daí as chamadas por request viram
# no-op: ALTER TABLE pega lock exclusivo mesmo com IF NOT EXISTS, e rodar isso
# a cada abertura do modal de tarefa disputava lock com o resto do sistema.
_FINANCEIRO_TABELAS_OK = False

def _ensure_financeiro_tables(cur):
    if _FINANCEIRO_TABELAS_OK:
        return
    cur.execute("""
        CREATE TABLE IF NOT EXISTS projetos (
            id SERIAL PRIMARY KEY,
            cliente_id INTEGER NOT NULL REFERENCES clientes(id),
            nome VARCHAR(100) NOT NULL,
            descricao TEXT DEFAULT '',
            orcamento NUMERIC(14,2) DEFAULT 0,
            ativo BOOLEAN DEFAULT TRUE,
            criado_em TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS lancamentos (
            id SERIAL PRIMARY KEY,
            projeto_id INTEGER NOT NULL REFERENCES projetos(id),
            descricao VARCHAR(200) NOT NULL,
            categoria VARCHAR(50) DEFAULT 'Outros',
            valor NUMERIC(14,2) NOT NULL,
            data_lancamento DATE DEFAULT CURRENT_DATE,
            criado_em TIMESTAMP DEFAULT NOW()
        )
    """)
    cur.execute("ALTER TABLE lancamentos ADD COLUMN IF NOT EXISTS localidade VARCHAR(150) DEFAULT ''")
    cur.execute("ALTER TABLE lancamentos ADD COLUMN IF NOT EXISTS tecnico VARCHAR(150) DEFAULT ''")
    cur.execute("ALTER TABLE projetos ADD COLUMN IF NOT EXISTS planilha_url TEXT DEFAULT ''")
    cur.execute("ALTER TABLE projetos ADD COLUMN IF NOT EXISTS planilha_mapeamento JSONB")
    cur.execute("ALTER TABLE projetos ADD COLUMN IF NOT EXISTS planilha_sync_em TIMESTAMP")
    cur.execute("ALTER TABLE projetos ADD COLUMN IF NOT EXISTS planilha_replace BOOLEAN DEFAULT FALSE")

class ProjetoModel(BaseModel):
    nome: str
    descricao: str = ""
    orcamento: float = 0.0
    responsavel_id: Optional[int] = None
    status_gestao: str = "EM ANDAMENTO"
    data_inicio: Optional[str] = None
    data_termino: Optional[str] = None
    escopo: str = ""

class LancamentoModel(BaseModel):
    descricao: str
    categoria: str = "Outros"
    valor: float
    data_lancamento: str = ""
    localidade: str = ""
    tecnico: str = ""

class LancamentoImportItem(BaseModel):
    descricao: str
    categoria: str = "Outros"
    valor: float
    data_lancamento: str = ""
    localidade: str = ""
    tecnico: str = ""

class ImportarProjetosPreviewBody(BaseModel):
    url: str

class ImportarProjetosBody(BaseModel):
    url: str
    col_cliente: str
    col_projeto: str
    col_orcamento: str = ""
    col_descricao: str = ""

class ImportarLancamentosBody(BaseModel):
    lancamentos: List[LancamentoImportItem]

class BaseImportacaoConfig(BaseModel):
    url: str
    col_cliente: str
    col_projeto: str
    col_orcamento: str = ""
    col_descricao: str = ""

@router.get("/api/config/base-importacao")
def get_base_importacao_config(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        import json as _json
        cur = conn.cursor()
        chave = f"base_importacao_{sess.get('time','Projetos')}"
        cur.execute("SELECT valor FROM configuracoes WHERE chave=%s", (chave,))
        row = cur.fetchone(); cur.close(); conn.close()
        if row:
            return _json.loads(row[0])
        return None
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.put("/api/config/base-importacao")
def salvar_base_importacao_config(body: BaseImportacaoConfig, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        import json as _json
        cur = conn.cursor()
        chave = f"base_importacao_{sess.get('time','Projetos')}"
        valor = _json.dumps({"url": body.url, "col_cliente": body.col_cliente,
                             "col_projeto": body.col_projeto, "col_orcamento": body.col_orcamento,
                             "col_descricao": body.col_descricao})
        cur.execute("""INSERT INTO configuracoes (chave, valor, atualizado_em) VALUES (%s, %s, NOW())
                       ON CONFLICT (chave) DO UPDATE SET valor=EXCLUDED.valor, atualizado_em=NOW()""", (chave, valor))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/base-importacao/atualizar")
async def atualizar_base_importacao(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox; openpyxl = _ox; _OPENPYXL_OK = True
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        import json as _json
        cur = conn.cursor()
        time_sess = sess.get("time", "Projetos")
        chave = f"base_importacao_{time_sess}"
        cur.execute("SELECT valor FROM configuracoes WHERE chave=%s", (chave,))
        row = cur.fetchone()
        if not row: raise HTTPException(status_code=400, detail="Base não configurada para este time. Configure a URL primeiro.")
        cfg = _json.loads(row[0])
        _ensure_financeiro_tables(cur)
        content = _download_planilha(cfg["url"])
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers, rows = _parse_sheet(ws)
        headers_lower = [h.lower() for h in headers]
        def col_idx(name):
            if not name: return None
            try: return headers_lower.index(name.lower())
            except ValueError: return None
        idx_cliente = col_idx(cfg["col_cliente"])
        idx_projeto = col_idx(cfg["col_projeto"])
        idx_orc     = col_idx(cfg.get("col_orcamento", ""))
        idx_desc    = col_idx(cfg.get("col_descricao", ""))
        if idx_cliente is None or idx_projeto is None:
            raise HTTPException(status_code=400, detail=f"Colunas não encontradas na planilha. Colunas disponíveis: {', '.join(headers)}. Reconfigure a base.")
        # Filtra clientes apenas do time do usuário logado
        cur.execute("SELECT id, LOWER(nome) FROM clientes WHERE ativo=TRUE AND COALESCE(time,'Projetos')=%s", (time_sess,))
        clientes_db = {r[1].strip(): r[0] for r in cur.fetchall()}
        criados = ignorados = ja_existem = 0
        clientes_nao_encontrados = set()
        for row in rows:
            c_nome = row[idx_cliente].strip() if idx_cliente < len(row) else ""
            p_nome = row[idx_projeto].strip() if idx_projeto < len(row) else ""
            if not c_nome or not p_nome: ignorados += 1; continue
            cid = clientes_db.get(c_nome.lower())
            if not cid: clientes_nao_encontrados.add(c_nome); ignorados += 1; continue
            orc = 0.0
            if idx_orc is not None and idx_orc < len(row):
                try:
                    v = row[idx_orc].replace('R$','').replace('.','').replace(',','.').strip()
                    orc = float(v)
                except Exception: pass
            desc = row[idx_desc].strip() if idx_desc is not None and idx_desc < len(row) else ""
            cur.execute("SELECT id FROM projetos WHERE cliente_id=%s AND LOWER(nome)=%s AND ativo=TRUE", (cid, p_nome.lower()))
            if cur.fetchone(): ja_existem += 1; continue
            cur.execute("INSERT INTO projetos (cliente_id, nome, descricao, orcamento) VALUES (%s,%s,%s,%s)", (cid, p_nome, desc, orc))
            criados += 1
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "criados": criados, "ja_existem": ja_existem, "ignorados": ignorados,
                "clientes_nao_encontrados": sorted(clientes_nao_encontrados)}
    except HTTPException: raise
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/importar-projetos/preview")
async def importar_projetos_preview(body: ImportarProjetosPreviewBody, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox; openpyxl = _ox; _OPENPYXL_OK = True
    try:
        content = _download_planilha(body.url)
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers, rows = _parse_sheet(ws)
        return {"headers": headers, "sample": rows[:5]}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.post("/api/importar-projetos/executar")
async def importar_projetos_executar(body: ImportarProjetosBody, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox; openpyxl = _ox; _OPENPYXL_OK = True
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        content = _download_planilha(body.url)
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers, rows = _parse_sheet(ws)

        headers_lower = [h.lower() for h in headers]
        def col_idx(name):
            if not name: return None
            try: return headers_lower.index(name.lower())
            except ValueError: return None

        idx_cliente  = col_idx(body.col_cliente)
        idx_projeto  = col_idx(body.col_projeto)
        idx_orc      = col_idx(body.col_orcamento) if body.col_orcamento else None
        idx_desc     = col_idx(body.col_descricao) if body.col_descricao else None

        if idx_cliente is None or idx_projeto is None:
            raise HTTPException(status_code=400, detail=f"Colunas não encontradas. Disponíveis: {', '.join(headers)}")

        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        time_sess = sess.get("time", "Projetos")
        cur.execute("SELECT id, LOWER(nome) FROM clientes WHERE ativo=TRUE AND COALESCE(time,'Projetos')=%s", (time_sess,))
        clientes_db = {r[1].strip(): r[0] for r in cur.fetchall()}

        criados = ignorados = ja_existem = 0
        clientes_nao_encontrados = set()
        for row in rows:
            c_nome = row[idx_cliente].strip() if idx_cliente < len(row) else ""
            p_nome = row[idx_projeto].strip() if idx_projeto < len(row) else ""
            if not c_nome or not p_nome:
                ignorados += 1; continue
            cid = clientes_db.get(c_nome.lower())
            if not cid:
                clientes_nao_encontrados.add(c_nome); ignorados += 1; continue
            orc = 0.0
            if idx_orc is not None and idx_orc < len(row):
                try:
                    v = row[idx_orc].replace('R$','').replace('.','').replace(',','.').strip()
                    orc = float(v)
                except Exception: pass
            desc = row[idx_desc].strip() if idx_desc is not None and idx_desc < len(row) else ""
            cur.execute("SELECT id FROM projetos WHERE cliente_id=%s AND LOWER(nome)=%s AND ativo=TRUE",
                        (cid, p_nome.lower()))
            if cur.fetchone():
                ja_existem += 1; continue
            cur.execute("INSERT INTO projetos (cliente_id, nome, descricao, orcamento) VALUES (%s,%s,%s,%s)",
                        (cid, p_nome, desc, orc))
            criados += 1

        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "criados": criados, "ja_existem": ja_existem, "ignorados": ignorados,
                "clientes_nao_encontrados": sorted(clientes_nao_encontrados)}
    except HTTPException: raise
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/opcoes-tarefa")
def opcoes_tarefa(faiston_token: str = Cookie(None)):
    """Retorna todos os clientes ativos (todos os times) + projetos ativos para
    uso no modal de tarefa.
    A lista de clientes continua deliberadamente sem filtro de time: os
    clientes existentes não têm o campo `time` confiável preenchido (o
    cadastro de cliente nunca teve tela própria até agora), então filtrar
    aqui hoje esvaziaria a lista pra Logística/Rede Credenciada. Escopar
    isso é seguro só depois de uma varredura confirmando o time de cada
    cliente existente -- registrado como próximo passo, não feito agora.
    A lista de projetos: com cliente, segue a mesma política do cliente
    (visível pra todo mundo, mesma ressalva acima); sem cliente (projeto
    interno, ex.: Desenvolvimento), só aparece pra quem é do mesmo time
    (ou admin, que vê tudo) -- aí o dado é confiável, porque é criado por
    criar_projeto_simples() e sempre carrega o próprio time."""
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401, detail="Não autenticado")
    conn = get_db()
    if not conn: raise HTTPException(status_code=500, detail="Banco offline")
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur); conn.commit()
        cur.execute("SELECT id, nome FROM clientes WHERE ativo=TRUE ORDER BY nome")
        clientes = [{"id": r[0], "nome": r[1]} for r in cur.fetchall()]
        tf = None if sess["perfil"] == "admin" else sess.get("time", "Projetos")
        cur.execute("""
            SELECT p.id, p.nome, COALESCE(c.nome, '')
            FROM projetos p
            LEFT JOIN clientes c ON c.id = p.cliente_id
            WHERE p.ativo = TRUE AND (c.id IS NULL OR c.ativo = TRUE)
              AND (p.cliente_id IS NOT NULL OR %s IS NULL OR COALESCE(p.time,'Projetos') = %s)
            ORDER BY c.nome NULLS FIRST, p.nome
        """, (tf, tf))
        projetos = [{"id": r[0], "nome": r[1], "cliente": r[2]} for r in cur.fetchall()]
        # Área sem projeto: o modal esconde o campo, e aqui a API não oferece o
        # que ela não pode usar. Admin atravessa áreas, então segue vendo tudo.
        if sess["perfil"] != "admin" and not _area_usa_projetos_cur(cur, sess.get("time", "Projetos")):
            projetos = []
        cur.close(); conn.close()
        return {"clientes": clientes, "projetos": projetos}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

class NovoProjetoSimples(BaseModel):
    nome: str
    descricao: str = ""
    time: str = ""
    cliente_id: Optional[int] = None
    orcamento: float = 0.0
    responsavel_id: Optional[int] = None
    status_gestao: str = "EM ANDAMENTO"
    data_inicio: Optional[str] = None
    data_termino: Optional[str] = None
    escopo: str = ""

@router.post("/api/projetos")
def criar_projeto_simples(p: NovoProjetoSimples, faiston_token: str = Cookie(None)):
    """Cadastro único de projeto -- cliente é opcional. Substitui, na tela de
    Gestão de Projetos e no quick-add de tarefa, o antigo caminho que exigia
    escolher um cliente antes de tudo. Com cliente_id: o projeto herda o
    time daquele cliente (mantém consistência -- nunca um projeto "do outro
    time" pendurado num cliente). Sem cliente_id: vira projeto interno,
    usando o time de quem cria (ou o escolhido, se for admin).

    POST /api/clientes/{id}/projetos continua existindo à parte -- é usado
    por financeiro.html, que já opera sempre dentro do contexto de um
    cliente específico."""
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    if not p.nome.strip(): raise HTTPException(status_code=400, detail="Informe o nome do projeto")
    status = p.status_gestao if p.status_gestao in ('EM ANDAMENTO', 'FINALIZAÇÃO', 'EM FREEZING') else 'EM ANDAMENTO'
    conn = get_db()
    if not conn: raise HTTPException(status_code=500, detail="Banco offline")
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        cliente_id = None
        if p.cliente_id:
            cur.execute("SELECT id, COALESCE(time,'Projetos') FROM clientes WHERE id=%s AND ativo=TRUE", (p.cliente_id,))
            row = cur.fetchone()
            if not row: raise HTTPException(status_code=400, detail="Cliente não encontrado")
            cliente_id, time_val = row
            if sess["perfil"] != "admin" and time_val != sess.get("time", "Projetos"):
                raise HTTPException(status_code=403, detail="Cliente não pertence ao seu time")
        else:
            time_val = p.time if (sess["perfil"] == "admin" and p.time in _times_validos_cur(cur)) else sess.get("time", "Projetos")
        _exigir_area_com_projeto(cur, time_val)
        cur.execute("""
            INSERT INTO projetos (cliente_id, nome, descricao, orcamento, escopo,
                                  responsavel_id, status_gestao, data_inicio, data_termino, time, ativo)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,TRUE) RETURNING id
        """, (cliente_id, p.nome.strip(), p.descricao, p.orcamento, p.escopo or p.descricao,
              p.responsavel_id, status, p.data_inicio or None, p.data_termino or None, time_val))
        new_id = cur.fetchone()[0]
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "id": new_id}
    except HTTPException: raise
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/todos-projetos")
def todos_projetos(faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur); conn.commit()
        cur.execute("""
            SELECT p.id, p.nome, c.nome
            FROM projetos p
            JOIN clientes c ON c.id = p.cliente_id
            WHERE p.ativo = TRUE AND c.ativo = TRUE
            ORDER BY c.nome, p.nome
        """)
        rows = cur.fetchall(); cur.close(); conn.close()
        return [{"id": r[0], "nome": r[1], "cliente": r[2]} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/projetos-by-cliente")
def projetos_by_cliente_nome(nome: str = "", faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        conn.commit()
        if nome:
            cur.execute("SELECT id FROM clientes WHERE nome=%s AND ativo=TRUE", (nome,))
            row = cur.fetchone()
            if not row: return []
            cid = row[0]
            cur.execute("SELECT id, nome FROM projetos WHERE cliente_id=%s AND ativo=TRUE ORDER BY nome", (cid,))
        else:
            cur.execute("SELECT id, nome FROM projetos WHERE ativo=TRUE ORDER BY nome")
        rows = cur.fetchall()
        cur.close(); conn.close()
        return [{"id": r[0], "nome": r[1]} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/clientes/{cid}/projetos")
def listar_projetos(cid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401)
    if sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        conn.commit()
        cur.execute("""
            SELECT p.id, p.nome, p.descricao, p.orcamento, p.criado_em,
                   COALESCE(SUM(l.valor), 0) AS gasto,
                   COALESCE(p.planilha_url, '') AS planilha_url,
                   p.planilha_mapeamento,
                   p.planilha_sync_em,
                   COALESCE(p.planilha_replace, FALSE) AS planilha_replace
            FROM projetos p
            LEFT JOIN lancamentos l ON l.projeto_id = p.id
            WHERE p.cliente_id = %s AND p.ativo = TRUE
            GROUP BY p.id ORDER BY p.criado_em DESC
        """, (cid,))
        rows = cur.fetchall()
        cur.close(); conn.close()
        return [{"id": r[0], "nome": r[1], "descricao": r[2],
                 "orcamento": float(r[3]), "criado_em": str(r[4])[:10],
                 "gasto": float(r[5]),
                 "planilha_url": r[6] or "",
                 "planilha_mapeamento": r[7] or {},
                 "planilha_sync_em": str(r[8])[:16] if r[8] else None,
                 "planilha_replace": bool(r[9])} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/clientes/{cid}/analise-financeira")
def analise_financeira(cid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401)
    if sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        conn.commit()

        # Gasto por localidade
        cur.execute("""
            SELECT COALESCE(NULLIF(TRIM(l.localidade),''), 'Não informado') AS loc,
                   SUM(l.valor) AS total
            FROM lancamentos l
            JOIN projetos p ON p.id = l.projeto_id
            WHERE p.cliente_id = %s AND p.ativo = TRUE
            GROUP BY loc ORDER BY total DESC LIMIT 10
        """, (cid,))
        loc_rows = cur.fetchall()
        total_loc = sum(r[1] for r in loc_rows) or 1
        por_localidade = [{"localidade": r[0], "gasto": float(r[1]),
                           "pct": round(float(r[1]) / total_loc * 100, 1)} for r in loc_rows]

        # Gasto por período (mês)
        cur.execute("""
            SELECT TO_CHAR(l.data_lancamento, 'YYYY-MM') AS periodo,
                   TO_CHAR(l.data_lancamento, 'Mon/YY') AS label,
                   SUM(l.valor) AS total
            FROM lancamentos l
            JOIN projetos p ON p.id = l.projeto_id
            WHERE p.cliente_id = %s AND p.ativo = TRUE
            GROUP BY periodo, label ORDER BY periodo
        """, (cid,))
        per_rows = cur.fetchall()
        max_per = max((float(r[2]) for r in per_rows), default=1)
        por_periodo = [{"periodo": r[0], "label": r[1], "gasto": float(r[2]),
                        "pct": round(float(r[2]) / max_per * 100, 1)} for r in per_rows]

        cur.close(); conn.close()
        return {"por_localidade": por_localidade, "por_periodo": por_periodo}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/clientes/{cid}/projetos")
def criar_projeto(cid: int, p: ProjetoModel, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        # O projeto criado aqui pertence à área do cliente -- se ela não trabalha
        # por projeto, esta porta também fica fechada.
        cur.execute("SELECT COALESCE(time,'Projetos') FROM clientes WHERE id=%s", (cid,))
        row_cli = cur.fetchone()
        if row_cli:
            _exigir_area_com_projeto(cur, row_cli[0])
        status = p.status_gestao if p.status_gestao in ('EM ANDAMENTO', 'FINALIZAÇÃO', 'EM FREEZING') else 'EM ANDAMENTO'
        cur.execute("""
            INSERT INTO projetos (cliente_id, nome, descricao, orcamento, escopo,
                                  responsavel_id, status_gestao, data_inicio, data_termino, ativo)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,TRUE) RETURNING id
        """, (cid, p.nome, p.descricao, p.orcamento, p.escopo or p.descricao,
              p.responsavel_id, status, p.data_inicio or None, p.data_termino or None))
        new_id = cur.fetchone()[0]
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "id": new_id}
    except HTTPException: raise
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.put("/api/projetos/{pid}")
def atualizar_projeto(pid: int, p: ProjetoModel, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        cur.execute("UPDATE projetos SET nome=%s, descricao=%s, orcamento=%s WHERE id=%s",
                    (p.nome, p.descricao, p.orcamento, pid))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.delete("/api/projetos/{pid}")
def deletar_projeto(pid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        cur.execute("UPDATE projetos SET ativo=FALSE WHERE id=%s", (pid,))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.get("/api/projetos/{pid}/lancamentos")
def listar_lancamentos(pid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess: raise HTTPException(status_code=401)
    if sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        cur.execute("""SELECT id, descricao, categoria, valor, data_lancamento, criado_em, localidade, tecnico
                       FROM lancamentos WHERE projeto_id=%s ORDER BY data_lancamento DESC, criado_em DESC""", (pid,))
        rows = cur.fetchall()
        cur.close(); conn.close()
        return [{"id": r[0], "descricao": r[1], "categoria": r[2],
                 "valor": float(r[3]), "data": str(r[4]), "criado_em": str(r[5])[:10],
                 "localidade": r[6] or "", "tecnico": r[7] or ""} for r in rows]
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.post("/api/projetos/{pid}/lancamentos")
def criar_lancamento(pid: int, l: LancamentoModel, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        data = l.data_lancamento or date.today().isoformat()
        cur.execute("INSERT INTO lancamentos (projeto_id, descricao, categoria, valor, data_lancamento, localidade, tecnico) VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id",
                    (pid, l.descricao, l.categoria, l.valor, data, l.localidade, l.tecnico))
        new_id = cur.fetchone()[0]
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "id": new_id}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

@router.delete("/api/lancamentos/{lid}")
def deletar_lancamento(lid: int, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM lancamentos WHERE id=%s", (lid,))
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))

def _detect_header_row(all_rows, max_search=10):
    """Return index of the row that best looks like a header (most non-empty text cells)."""
    best_idx, best_score = 0, 0
    for i, row in enumerate(all_rows[:max_search]):
        score = sum(1 for c in row if c is not None and isinstance(c, str) and c.strip())
        if score > best_score:
            best_score, best_idx = score, i
    return best_idx

def _parse_sheet(ws):
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return [], []
    hi = _detect_header_row(all_rows)
    headers = [str(h).strip() if h is not None else "" for h in all_rows[hi]]
    rows = []
    for row in all_rows[hi+1:]:
        if not any(c is not None and str(c).strip() for c in row):
            continue
        rows.append([str(c).strip() if c is not None else "" for c in row])
    return headers, rows

@router.post("/api/projetos/{pid}/parse-planilha")
async def parse_planilha(pid: int, file: UploadFile = File(...),
                          sheet_name: str = "", faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    global _OPENPYXL_OK, openpyxl
    if not _OPENPYXL_OK:
        import subprocess, sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
        import openpyxl as _ox
        openpyxl = _ox
        _OPENPYXL_OK = True
    try:
        content = await file.read()
        filename = (file.filename or "").lower()

        if filename.endswith(".csv"):
            text = content.decode("utf-8-sig", errors="replace")
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except Exception:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect)
            all_rows = list(reader)
            if not all_rows:
                return {"sheets": [], "selected_sheet": "", "headers": [], "rows": []}
            hi = _detect_header_row([[c for c in r] for r in all_rows])
            headers = [str(h).strip() for h in all_rows[hi]]
            rows = [[str(c).strip() for c in r]
                    for r in all_rows[hi+1:] if any(c.strip() for c in r)]
            return {"sheets": [], "selected_sheet": filename, "headers": headers, "rows": rows[:3000]}

        # Excel
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

        # Discover sheets that have useful data (≥4 non-empty header cells)
        usable = []
        for sname in wb.sheetnames:
            ws = wb[sname]
            rows_peek = list(ws.iter_rows(min_row=1, max_row=10, values_only=True))
            hi = _detect_header_row(rows_peek)
            hrow = rows_peek[hi] if rows_peek else []
            n_text = sum(1 for c in hrow if c and isinstance(c, str) and c.strip())
            if n_text >= 4:
                usable.append(sname)

        if not usable:
            raise ValueError("Nenhuma aba com dados tabulares encontrada")

        # Pick sheet
        chosen = sheet_name if sheet_name in usable else usable[0]
        # Prefer sheets with "VALOR FINAL"
        if not sheet_name:
            for s in usable:
                ws_tmp = wb[s]
                rows_tmp = list(ws_tmp.iter_rows(min_row=1, max_row=5, values_only=True))
                hi_tmp = _detect_header_row(rows_tmp)
                hrow_tmp = rows_tmp[hi_tmp] if rows_tmp else []
                if any('VALOR FINAL' in str(c) for c in hrow_tmp if c):
                    chosen = s
                    break

        ws = wb[chosen]
        headers, rows = _parse_sheet(ws)
        return {"sheets": usable, "selected_sheet": chosen, "headers": headers, "rows": rows[:3000]}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Erro ao processar arquivo: {str(e)}")

@router.post("/api/projetos/{pid}/importar-lancamentos")
def importar_lancamentos(pid: int, body: ImportarLancamentosBody, faiston_token: str = Cookie(None)):
    sess = get_session(faiston_token)
    if not sess or sess["perfil"] not in ("admin", "gestor", "demo"): raise HTTPException(status_code=403)
    conn = get_db()
    if not conn: raise HTTPException(status_code=500)
    try:
        cur = conn.cursor()
        _ensure_financeiro_tables(cur)
        today = date.today().isoformat()
        count = 0
        for l in body.lancamentos:
            data = l.data_lancamento or today
            cur.execute(
                "INSERT INTO lancamentos (projeto_id, descricao, categoria, valor, data_lancamento, localidade, tecnico) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                (pid, l.descricao[:200], l.categoria, l.valor, data, l.localidade[:150], l.tecnico[:150])
            )
            count += 1
        conn.commit(); cur.close(); conn.close()
        return {"sucesso": True, "importados": count}
    except Exception as e: raise HTTPException(status_code=500, detail=str(e))


def _resolve_download_url(url: str) -> str:
    """Convert OneDrive/SharePoint share links to direct download URLs."""
    import urllib.parse
    u = url.strip()
    # SharePoint / OneDrive for Business
    if 'sharepoint.com' in u or 'onedrive.live.com' in u:
        sep = '&' if '?' in u else '?'
        return u + sep + 'download=1'
    # 1drv.ms short link — follow redirect then add download=1
    if '1drv.ms' in u or 'onedrive.com' in u:
        try:
            import requests as _req
            r = _req.head(u, timeout=15, allow_redirects=True)
            resolved = r.url
            sep = '&' if '?' in resolved else '?'
            return resolved + sep + 'download=1'
        except Exception:
            pass
    return u


def _download_planilha(url: str) -> bytes:
    import requests as _req
    dl_url = _resolve_download_url(url)
    r = _req.get(dl_url, timeout=60, allow_redirects=True,
                 headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code != 200:
        raise ValueError(f"Erro ao baixar planilha (HTTP {r.status_code}). Verifique se o link permite acesso sem login.")
    if len(r.content) < 50:
        raise ValueError("Arquivo baixado está vazio. Verifique as permissões do link de compartilhamento.")
    return r.content
