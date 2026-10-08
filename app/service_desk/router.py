"""Endpoints do Service Desk (/service-desk e /api/sd/*).

Premissa do módulo: registrar um atendimento tem que ser mais rápido que
anotar num bloco de notas. Por isso:
- /api/sd/opcoes devolve tudo que a tela precisa num request só (filas,
  canais, categorias, status, modelos rápidos tirados do histórico);
- fila/categoria/canal novos digitados no registro são criados na hora,
  sem precisar de cadastro prévio;
- redirecionar pra outra fila já marca o status como "redirecionado".

Acesso: time 'Service Desk' ou cargo sd_*; admin/diretor atravessam.
Operador vê só os próprios atendimentos; supervisor vê a equipe toda.
"""
import io
from datetime import date, datetime, time, timedelta
from typing import List, Optional

from fastapi import APIRouter, Cookie, HTTPException, Query
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel, Field

from app.service_desk import db
from app.service_desk.turno import (STATUS_OPERADOR, agora, deve_encerrar, dias_12x36, escolher_escala,
                                    janela_contagem, janela_turno)

router = APIRouter()

TIME_SERVICE_DESK = "Service Desk"
CARGOS_SD = ("sd_operador", "sd_supervisor")
TABELAS_OPCAO = {"filas": "sd_filas", "canais": "sd_canais", "categorias": "sd_categorias"}
TIPOS_ESCALA = ("turno", "plantao", "sobreaviso")


# ── Acesso ───────────────────────────────────────────────────────────────
def eh_admin(sess: dict) -> bool:
    return sess.get("perfil") in ("admin", "diretor")


def pode_usar(sess: Optional[dict]) -> bool:
    if not sess:
        return False
    return eh_admin(sess) or sess.get("time") == TIME_SERVICE_DESK or sess.get("cargo") in CARGOS_SD


def eh_supervisor(sess: dict) -> bool:
    if eh_admin(sess) or sess.get("cargo") == "sd_supervisor":
        return True
    return sess.get("perfil") == "gestor" and sess.get("time") == TIME_SERVICE_DESK


def _sessao(token: Optional[str]) -> dict:
    sess = db.get_session(token)
    if not sess:
        raise HTTPException(status_code=401, detail="Não autenticado")
    if not pode_usar(sess):
        raise HTTPException(status_code=403, detail="Service Desk é restrito à equipe da operação")
    return sess


def _supervisor(token: Optional[str]) -> dict:
    sess = _sessao(token)
    if not eh_supervisor(sess):
        raise HTTPException(status_code=403, detail="Só supervisor do Service Desk")
    return sess


def _conn():
    conn = db.get_conn()
    if not conn:
        raise HTTPException(status_code=500, detail="Banco offline")
    return conn


def _hora(v: Optional[str]) -> Optional[time]:
    if not v:
        return None
    try:
        return time.fromisoformat(v)
    except ValueError:
        raise HTTPException(status_code=400, detail="Horário inválido — use HH:MM")


def _data(v: Optional[str], campo: str = "data") -> Optional[date]:
    if not v:
        return None
    try:
        return date.fromisoformat(v)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{campo} inválida — use AAAA-MM-DD")


def _hhmm(t: Optional[time]) -> str:
    return t.strftime("%H:%M") if t else ""


# ── Página ───────────────────────────────────────────────────────────────
@router.get("/service-desk")
def pagina(faiston_token: str = Cookie(None)):
    sess = db.get_session(faiston_token)
    if not sess:
        return RedirectResponse("/")
    if not pode_usar(sess):
        return RedirectResponse("/funcionario")
    return FileResponse("static/service-desk.html", headers={"Cache-Control": "no-cache"})


# ── Operador / jornada ───────────────────────────────────────────────────
def _operador(cur, uid: int) -> dict:
    cur.execute("""
        SELECT u.nome, COALESCE(o.matricula,''), COALESCE(o.nivel,'N1'), COALESCE(o.site,''),
               o.jornada_inicio, o.jornada_fim, COALESCE(o.regime,''), COALESCE(o.meta_turno, 20)
        FROM usuarios u LEFT JOIN sd_operadores o ON o.usuario_id = u.id
        WHERE u.id = %s
    """, (uid,))
    r = cur.fetchone()
    if not r:
        raise HTTPException(status_code=404, detail="Operador não encontrado")
    return {"id": uid, "nome": r[0], "matricula": r[1], "nivel": r[2], "site": r[3],
            "jornada_inicio": r[4], "jornada_fim": r[5], "regime": r[6], "meta_turno": r[7]}


def _jornada_em(cur, op: dict, ref: datetime):
    """Jornada do turno ao qual `ref` pertence: escala cadastrada manda (ver
    escolher_escala -- olha ontem também, às 02h o plantão em curso é o que
    começou ontem); sem escala vale a jornada padrão."""
    cur.execute("""
        SELECT data, hora_inicio, hora_fim, tipo FROM sd_escalas
        WHERE usuario_id = %s AND data IN (%s, %s)
    """, (op["id"], ref.date(), ref.date() - timedelta(days=1)))
    escala = escolher_escala(cur.fetchall(), ref)
    if escala:
        return escala[1], escala[2], escala[3]
    return op["jornada_inicio"], op["jornada_fim"], None


def _encerrar_esquecidos(cur, uid: Optional[int] = None):
    """Fecha status aberto que passou do fim do turno (a pessoa saiu em
    pausa e não voltou). Roda no job e também antes de ler o status.
    O turno é o da escala do dia do evento, não só a jornada padrão --
    senão o plantão noturno de quem é 08h-17h no cadastro era fechado no meio."""
    sql = """
        SELECT e.id, e.inicio, e.usuario_id, o.jornada_inicio, o.jornada_fim
        FROM sd_status_eventos e LEFT JOIN sd_operadores o ON o.usuario_id = e.usuario_id
        WHERE e.fim IS NULL
    """
    params: tuple = ()
    if uid is not None:
        sql += " AND e.usuario_id = %s"
        params = (uid,)
    cur.execute(sql, params)
    ref = agora()
    fechados = 0
    for eid, inicio, dono, jp_ini, jp_fim in cur.fetchall():
        op = {"id": dono, "jornada_inicio": jp_ini, "jornada_fim": jp_fim}
        j_ini, j_fim, _ = _jornada_em(cur, op, inicio)
        fim = deve_encerrar(inicio, j_ini, j_fim, ref)
        if fim:
            cur.execute("UPDATE sd_status_eventos SET fim = %s, auto_encerrado = TRUE WHERE id = %s",
                        (max(fim, inicio), eid))
            fechados += 1
    return fechados


def job_encerrar_status_esquecidos():
    conn = db.get_conn()
    if not conn:
        return
    try:
        cur = conn.cursor()
        n = _encerrar_esquecidos(cur)
        conn.commit(); cur.close()
        if n:
            print(f"[service_desk] {n} status esquecido(s) encerrado(s) no fim do turno")
    except Exception as e:
        print(f"[service_desk] Erro ao encerrar status esquecidos: {e}")
    finally:
        conn.close()


def _status_atual(cur, uid: int) -> dict:
    cur.execute("""
        SELECT status, motivo, inicio FROM sd_status_eventos
        WHERE usuario_id = %s AND fim IS NULL ORDER BY inicio DESC LIMIT 1
    """, (uid,))
    r = cur.fetchone()
    if not r:
        return {"status": "offline", "motivo": "", "desde": None, "segundos": 0}
    return {"status": r[0], "motivo": r[1] or "", "desde": r[2].isoformat(),
            "segundos": max(0, int((agora() - r[2]).total_seconds()))}


def _kpis(cur, uid: Optional[int], ini: datetime, fim: datetime, meta: int) -> dict:
    filtro_u = "AND a.usuario_id = %s" if uid else ""
    params = (ini, fim) + ((uid,) if uid else ())
    cur.execute(f"""
        SELECT COUNT(*),
               COUNT(*) FILTER (WHERE t.finaliza),
               COUNT(*) FILTER (WHERE a.status = 'concluido' AND COALESCE(a.fila_destino,'') = ''),
               COUNT(*) FILTER (WHERE t.finaliza AND NOT (a.status = 'concluido' AND COALESCE(a.fila_destino,'') = ''))
        FROM sd_atendimentos a LEFT JOIN sd_status_tipos t ON t.chave = a.status
        WHERE a.criado_em >= %s AND a.criado_em < %s {filtro_u}
    """, params)
    total, finalizados, resolvidos, outros = cur.fetchone()
    cur.execute(f"""
        SELECT a.fila_entrada, COUNT(*) FROM sd_atendimentos a
        WHERE a.criado_em >= %s AND a.criado_em < %s {filtro_u}
        GROUP BY a.fila_entrada ORDER BY COUNT(*) DESC LIMIT 1
    """, params)
    top = cur.fetchone()
    return {
        "total": total, "meta": meta,
        "fcr_pct": round(100 * resolvidos / finalizados) if finalizados else None,
        "resolvidos_1o_contato": resolvidos, "escalados_outros": outros,
        "fila_top": top[0] if top else None,
        "fila_top_pct": round(100 * top[1] / total) if top and total else None,
    }


@router.get("/api/sd/me")
def me(faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        _encerrar_esquecidos(cur, sess["id"])
        conn.commit()
        op = _operador(cur, sess["id"])
        ref = agora()
        j_ini, j_fim, tipo_escala = _jornada_em(cur, op, ref)
        ini, fim = janela_turno(ref, j_ini, j_fim)
        kpis = _kpis(cur, sess["id"], *janela_contagem(ref, j_ini, j_fim), op["meta_turno"])
        status = _status_atual(cur, sess["id"])
        cur.close()
        return {
            "operador": {**op, "jornada_inicio": _hhmm(op["jornada_inicio"]), "jornada_fim": _hhmm(op["jornada_fim"])},
            "supervisor": eh_supervisor(sess),
            "status": status,
            "turno": {"inicio": ini.isoformat(), "fim": fim.isoformat(),
                      "hora_inicio": _hhmm(j_ini), "hora_fim": _hhmm(j_fim), "tipo": tipo_escala},
            "kpis": kpis,
        }
    finally:
        conn.close()


class StatusModel(BaseModel):
    status: str
    motivo: str = ""


@router.post("/api/sd/status")
def trocar_status(body: StatusModel, faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    if body.status not in STATUS_OPERADOR + ("offline",):
        raise HTTPException(status_code=400, detail="Status inválido")
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("UPDATE sd_status_eventos SET fim = NOW() WHERE usuario_id = %s AND fim IS NULL", (sess["id"],))
        if body.status != "offline":
            cur.execute("INSERT INTO sd_status_eventos (usuario_id, status, motivo) VALUES (%s, %s, %s)",
                        (sess["id"], body.status, (body.motivo or "")[:60]))
        conn.commit()
        st = _status_atual(cur, sess["id"])
        cur.close()
        return st
    finally:
        conn.close()


# ── Opções (filas, canais, categorias, status, modelos rápidos) ──────────
@router.get("/api/sd/opcoes")
def opcoes(faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        out = {}
        for chave, tabela in TABELAS_OPCAO.items():
            # Mais usadas primeiro -- o que a pessoa mais escolhe fica no topo.
            coluna = {"filas": "fila_entrada", "canais": "canal", "categorias": "categoria"}[chave]
            cur.execute(f"""
                SELECT o.nome FROM {tabela} o
                LEFT JOIN (SELECT {coluna} AS nome, COUNT(*) AS n FROM sd_atendimentos
                           WHERE criado_em > NOW() - INTERVAL '60 days' GROUP BY {coluna}) u ON u.nome = o.nome
                WHERE o.ativo ORDER BY COALESCE(u.n, 0) DESC, o.nome
            """)
            out[chave] = [r[0] for r in cur.fetchall()]
        cur.execute("SELECT chave, rotulo, cor, finaliza FROM sd_status_tipos WHERE ativo ORDER BY ordem, rotulo")
        out["status"] = [{"chave": r[0], "rotulo": r[1], "cor": r[2], "finaliza": r[3]} for r in cur.fetchall()]
        # Modelos rápidos: combinações categoria+problema+tratativa que a
        # própria pessoa mais repete. Sem cadastro -- sai do histórico.
        cur.execute("""
            SELECT categoria, problema, tratativa, COUNT(*) AS n, MAX(criado_em)
            FROM sd_atendimentos
            WHERE usuario_id = %s AND criado_em > NOW() - INTERVAL '60 days'
            GROUP BY categoria, problema, tratativa
            HAVING COUNT(*) >= 2
            ORDER BY n DESC, MAX(criado_em) DESC LIMIT 8
        """, (sess["id"],))
        out["modelos"] = [{"categoria": r[0], "problema": r[1], "tratativa": r[2], "usos": r[3]} for r in cur.fetchall()]
        # Último canal/fila usados -- a tela já abre com eles escolhidos.
        cur.execute("""
            SELECT canal, fila_entrada FROM sd_atendimentos WHERE usuario_id = %s
            ORDER BY criado_em DESC LIMIT 1
        """, (sess["id"],))
        r = cur.fetchone()
        out["ultimo"] = {"canal": r[0], "fila_entrada": r[1]} if r else None
        cur.close()
        return out
    finally:
        conn.close()


class OpcaoModel(BaseModel):
    nome: str = Field(min_length=1, max_length=120)


def _garantir_opcao(cur, tabela: str, nome: str):
    nome = (nome or "").strip()
    if nome:
        cur.execute(f"INSERT INTO {tabela} (nome) VALUES (%s) ON CONFLICT (nome) DO UPDATE SET ativo = TRUE", (nome,))


@router.post("/api/sd/opcoes/{tipo}")
def criar_opcao(tipo: str, body: OpcaoModel, faiston_token: str = Cookie(None)):
    _sessao(faiston_token)
    tabela = TABELAS_OPCAO.get(tipo)
    if not tabela:
        raise HTTPException(status_code=404, detail="Tipo de opção inválido")
    conn = _conn()
    try:
        cur = conn.cursor()
        _garantir_opcao(cur, tabela, body.nome)
        conn.commit(); cur.close()
        return {"sucesso": True, "nome": body.nome.strip()}
    finally:
        conn.close()


@router.delete("/api/sd/opcoes/{tipo}/{nome}")
def desativar_opcao(tipo: str, nome: str, faiston_token: str = Cookie(None)):
    """Tira da lista sem apagar -- atendimento antigo continua com o nome."""
    _supervisor(faiston_token)
    tabela = TABELAS_OPCAO.get(tipo)
    if not tabela:
        raise HTTPException(status_code=404, detail="Tipo de opção inválido")
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute(f"UPDATE {tabela} SET ativo = FALSE WHERE nome = %s", (nome,))
        conn.commit(); cur.close()
        return {"sucesso": True}
    finally:
        conn.close()


class StatusTipoModel(BaseModel):
    rotulo: str = Field(min_length=1, max_length=60)
    cor: str = "slate"
    finaliza: bool = False


@router.post("/api/sd/status-tipos")
def criar_status_tipo(body: StatusTipoModel, faiston_token: str = Cookie(None)):
    _supervisor(faiston_token)
    import re
    import unicodedata
    base = unicodedata.normalize("NFKD", body.rotulo).encode("ascii", "ignore").decode().lower()
    chave = re.sub(r"[^a-z0-9]+", "_", base).strip("_")[:30]
    if not chave:
        raise HTTPException(status_code=400, detail="Nome de status inválido")
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO sd_status_tipos (chave, rotulo, cor, finaliza)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (chave) DO UPDATE SET rotulo = EXCLUDED.rotulo, cor = EXCLUDED.cor,
                                              finaliza = EXCLUDED.finaliza, ativo = TRUE
        """, (chave, body.rotulo.strip(), body.cor[:20], body.finaliza))
        conn.commit(); cur.close()
        return {"sucesso": True, "chave": chave}
    finally:
        conn.close()


@router.get("/api/sd/solicitantes")
def solicitantes(q: str = "", faiston_token: str = Cookie(None)):
    """Autocompletar do nome do solicitante -- devolve o login junto pra
    preencher o segundo campo sozinho. Restrito à equipe SD (LGPD)."""
    _sessao(faiston_token)
    q = (q or "").strip()
    if len(q) < 2:
        return []
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT solicitante_nome, solicitante_login, MAX(criado_em) FROM sd_atendimentos
            WHERE solicitante_nome ILIKE %s OR solicitante_login ILIKE %s
            GROUP BY solicitante_nome, solicitante_login ORDER BY MAX(criado_em) DESC LIMIT 8
        """, (f"%{q}%", f"%{q}%"))
        out = [{"nome": r[0], "login": r[1] or ""} for r in cur.fetchall()]
        cur.close()
        return out
    finally:
        conn.close()


# ── Atendimentos ─────────────────────────────────────────────────────────
class AtendimentoModel(BaseModel):
    solicitante_nome: str = Field(min_length=1, max_length=150)
    solicitante_login: str = Field(default="", max_length=80)
    canal: str = Field(default="", max_length=60)
    fila_entrada: str = Field(min_length=1, max_length=120)
    fila_destino: str = Field(default="", max_length=120)
    categoria: str = Field(default="", max_length=120)
    problema: str = Field(min_length=1, max_length=300)
    tratativa: str = Field(min_length=1)
    status: str = ""
    chamado_externo: str = Field(default="", max_length=60)


def _normalizar(body: AtendimentoModel, cur) -> dict:
    d = {k: (v.strip() if isinstance(v, str) else v) for k, v in body.model_dump().items()}
    for campo in ("solicitante_nome", "fila_entrada", "problema", "tratativa"):
        if not d[campo]:
            raise HTTPException(status_code=400, detail=f"Campo obrigatório: {campo}")
    if d["fila_destino"] == d["fila_entrada"]:
        d["fila_destino"] = ""
    if not d["status"]:
        d["status"] = "redirecionado" if d["fila_destino"] else "concluido"
    cur.execute("SELECT 1 FROM sd_status_tipos WHERE chave = %s AND ativo", (d["status"],))
    if not cur.fetchone():
        raise HTTPException(status_code=400, detail="Status inválido")
    _garantir_opcao(cur, "sd_filas", d["fila_entrada"])
    _garantir_opcao(cur, "sd_filas", d["fila_destino"])
    _garantir_opcao(cur, "sd_canais", d["canal"])
    _garantir_opcao(cur, "sd_categorias", d["categoria"])
    return d


_COLUNAS = ("a.id, a.usuario_id, u.nome, a.solicitante_nome, a.solicitante_login, a.canal, a.fila_entrada, "
            "a.fila_destino, a.categoria, a.problema, a.tratativa, a.status, a.chamado_externo, a.criado_em, "
            "a.atualizado_em")


def _linha(r) -> dict:
    return {"id": r[0], "usuario_id": r[1], "operador": r[2], "solicitante_nome": r[3], "solicitante_login": r[4],
            "canal": r[5], "fila_entrada": r[6], "fila_destino": r[7], "categoria": r[8], "problema": r[9],
            "tratativa": r[10], "status": r[11], "chamado_externo": r[12], "criado_em": r[13].isoformat(),
            "atualizado_em": r[14].isoformat() if r[14] else None}


def _buscar(cur, aid: int) -> dict:
    cur.execute(f"SELECT {_COLUNAS} FROM sd_atendimentos a JOIN usuarios u ON u.id = a.usuario_id WHERE a.id = %s", (aid,))
    r = cur.fetchone()
    if not r:
        raise HTTPException(status_code=404, detail="Atendimento não encontrado")
    return _linha(r)


@router.post("/api/sd/atendimentos")
def criar_atendimento(body: AtendimentoModel, faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        d = _normalizar(body, cur)
        cur.execute("""
            INSERT INTO sd_atendimentos (usuario_id, solicitante_nome, solicitante_login, canal, fila_entrada,
                fila_destino, categoria, problema, tratativa, status, chamado_externo)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id
        """, (sess["id"], d["solicitante_nome"], d["solicitante_login"], d["canal"], d["fila_entrada"],
              d["fila_destino"], d["categoria"], d["problema"], d["tratativa"], d["status"], d["chamado_externo"]))
        aid = cur.fetchone()[0]
        conn.commit()
        out = _buscar(cur, aid)
        cur.close()
        return out
    finally:
        conn.close()


def _filtros_lista(sess, escopo, status, fila, q, operador, cur):
    where, params = [], []
    alvo = None  # None = equipe toda (supervisor sem filtro de operador)
    if eh_supervisor(sess) and operador != "eu":
        if operador:
            if not operador.isdigit():
                raise HTTPException(status_code=400, detail="Operador inválido")
            alvo = int(operador)
    else:
        alvo = sess["id"]
    if alvo is not None:
        where.append("a.usuario_id = %s"); params.append(alvo)
    if escopo == "turno":
        # Cada operador na janela do PRÓPRIO turno -- antes valia a de quem
        # estava olhando, e o supervisor 08h-17h via o plantão 19h-07h cortado.
        ref = agora()

        def janela(uid):
            j_ini, j_fim, _ = _jornada_em(cur, _operador(cur, uid), ref)
            return janela_contagem(ref, j_ini, j_fim)

        if alvo is not None:
            where.append("a.criado_em >= %s AND a.criado_em < %s"); params += list(janela(alvo))
        else:
            equipe = _ids_equipe(cur)
            partes = []
            for uid in equipe:
                partes.append("(a.usuario_id = %s AND a.criado_em >= %s AND a.criado_em < %s)")
                params += [uid, *janela(uid)]
            # Quem não é da equipe (ex.: admin que registrou) segue na janela de quem olha.
            ini, fim = janela(sess["id"])
            partes.append("(NOT a.usuario_id = ANY(%s) AND a.criado_em >= %s AND a.criado_em < %s)")
            params += [equipe, ini, fim]
            where.append("(" + " OR ".join(partes) + ")")
    elif escopo == "hoje":
        where.append("a.criado_em::date = CURRENT_DATE")
    elif escopo == "7d":
        where.append("a.criado_em > NOW() - INTERVAL '7 days'")
    elif escopo == "30d":
        where.append("a.criado_em > NOW() - INTERVAL '30 days'")
    if status:
        where.append("a.status = %s"); params.append(status)
    if fila:
        where.append("(a.fila_entrada = %s OR a.fila_destino = %s)"); params += [fila, fila]
    if q:
        where.append("""(a.solicitante_nome ILIKE %s OR a.solicitante_login ILIKE %s OR a.problema ILIKE %s
                         OR a.tratativa ILIKE %s OR a.categoria ILIKE %s OR a.chamado_externo ILIKE %s)""")
        params += [f"%{q}%"] * 6
    return (" WHERE " + " AND ".join(where)) if where else "", params


@router.get("/api/sd/atendimentos")
def listar_atendimentos(escopo: str = "turno", status: str = "", fila: str = "", q: str = "",
                        operador: str = "", limite: int = Query(200, le=1000),
                        faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        where, params = _filtros_lista(sess, escopo, status, fila, q.strip(), operador, cur)
        cur.execute(f"""
            SELECT {_COLUNAS} FROM sd_atendimentos a JOIN usuarios u ON u.id = a.usuario_id
            {where} ORDER BY a.criado_em DESC LIMIT %s
        """, params + [limite])
        out = [_linha(r) for r in cur.fetchall()]
        cur.close()
        return out
    finally:
        conn.close()


def _pode_editar(sess, atend) -> bool:
    return eh_supervisor(sess) or atend["usuario_id"] == sess["id"]


@router.put("/api/sd/atendimentos/{aid}")
def editar_atendimento(aid: int, body: AtendimentoModel, faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        if not _pode_editar(sess, _buscar(cur, aid)):
            raise HTTPException(status_code=403, detail="Só dá pra editar os próprios atendimentos")
        d = _normalizar(body, cur)
        cur.execute("""
            UPDATE sd_atendimentos SET solicitante_nome=%s, solicitante_login=%s, canal=%s, fila_entrada=%s,
                fila_destino=%s, categoria=%s, problema=%s, tratativa=%s, status=%s, chamado_externo=%s,
                atualizado_em=NOW()
            WHERE id=%s
        """, (d["solicitante_nome"], d["solicitante_login"], d["canal"], d["fila_entrada"], d["fila_destino"],
              d["categoria"], d["problema"], d["tratativa"], d["status"], d["chamado_externo"], aid))
        conn.commit()
        out = _buscar(cur, aid)
        cur.close()
        return out
    finally:
        conn.close()


@router.delete("/api/sd/atendimentos/{aid}")
def excluir_atendimento(aid: int, faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        if not _pode_editar(sess, _buscar(cur, aid)):
            raise HTTPException(status_code=403, detail="Só dá pra excluir os próprios atendimentos")
        cur.execute("DELETE FROM sd_atendimentos WHERE id = %s", (aid,))
        conn.commit(); cur.close()
        return {"sucesso": True}
    finally:
        conn.close()


@router.get("/api/sd/atendimentos/exportar")
def exportar_atendimentos(escopo: str = "30d", status: str = "", fila: str = "", q: str = "",
                          operador: str = "", faiston_token: str = Cookie(None)):
    sess = _sessao(faiston_token)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl não instalado")
    conn = _conn()
    try:
        cur = conn.cursor()
        where, params = _filtros_lista(sess, escopo, status, fila, q.strip(), operador, cur)
        cur.execute(f"""
            SELECT {_COLUNAS} FROM sd_atendimentos a JOIN usuarios u ON u.id = a.usuario_id
            {where} ORDER BY a.criado_em DESC LIMIT 20000
        """, params)
        linhas = [_linha(r) for r in cur.fetchall()]
        cur.execute("SELECT chave, rotulo FROM sd_status_tipos")
        rotulos = dict(cur.fetchall())
        cur.close()
    finally:
        conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Atendimentos"
    cab = ["Data", "Hora", "Operador", "Solicitante", "Login", "Canal", "Fila de entrada", "Redirecionado para",
           "Categoria", "Problema", "Tratativa", "Status", "Chamado externo"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B0D1F")
    for a in linhas:
        dt = datetime.fromisoformat(a["criado_em"])
        ws.append([dt.strftime("%d/%m/%Y"), dt.strftime("%H:%M"), a["operador"], a["solicitante_nome"],
                   a["solicitante_login"], a["canal"], a["fila_entrada"], a["fila_destino"], a["categoria"],
                   a["problema"], a["tratativa"], rotulos.get(a["status"], a["status"]), a["chamado_externo"]])
    for col, larg in zip("ABCDEFGHIJKLM", (11, 7, 22, 26, 16, 12, 30, 30, 18, 36, 50, 16, 16)):
        ws.column_dimensions[col].width = larg
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    nome = f"atendimentos_service_desk_{date.today().strftime('%d-%m-%Y')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nome}"'})


# ── Visão gerencial (dashboard do supervisor) ────────────────────────────
PERIODOS_DASHBOARD = {"hoje": 1, "7d": 7, "30d": 30, "90d": 90}


@router.get("/api/sd/dashboard")
def dashboard(periodo: str = "7d", faiston_token: str = Cookie(None)):
    """Métricas agregadas da operação pro supervisor (aba Service Desk do
    /dashboard). Período em dias corridos terminando hoje."""
    _supervisor(faiston_token)
    dias = PERIODOS_DASHBOARD.get(periodo)
    if not dias:
        raise HTTPException(status_code=400, detail="Período inválido")
    hoje = agora().date()
    d0 = hoje - timedelta(days=dias - 1)
    ini, fim = datetime.combine(d0, time(0)), datetime.combine(hoje + timedelta(days=1), time(0))
    conn = _conn()
    try:
        cur = conn.cursor()
        _encerrar_esquecidos(cur)
        conn.commit()
        janela = "a.criado_em >= %s AND a.criado_em < %s"
        resolvido = "(a.status = 'concluido' AND COALESCE(a.fila_destino,'') = '')"
        cur.execute(f"""
            SELECT COUNT(*),
                   COUNT(*) FILTER (WHERE t.finaliza),
                   COUNT(*) FILTER (WHERE {resolvido}),
                   COUNT(*) FILTER (WHERE COALESCE(a.fila_destino,'') <> ''),
                   COUNT(*) FILTER (WHERE NOT COALESCE(t.finaliza, FALSE)),
                   COUNT(DISTINCT a.usuario_id)
            FROM sd_atendimentos a LEFT JOIN sd_status_tipos t ON t.chave = a.status
            WHERE {janela}
        """, (ini, fim))
        total, finalizados, resolvidos, redirecionados, em_aberto, operadores = cur.fetchone()

        cur.execute(f"""
            SELECT a.criado_em::date, COUNT(*), COUNT(*) FILTER (WHERE {resolvido})
            FROM sd_atendimentos a WHERE {janela} GROUP BY 1 ORDER BY 1
        """, (ini, fim))
        por_data = {r[0]: (r[1], r[2]) for r in cur.fetchall()}
        por_dia = []
        for i in range(dias):
            d = d0 + timedelta(days=i)
            t, r = por_data.get(d, (0, 0))
            por_dia.append({"data": d.isoformat(), "total": t, "resolvidos": r})

        cur.execute(f"""
            SELECT EXTRACT(HOUR FROM a.criado_em)::int, COUNT(*)
            FROM sd_atendimentos a WHERE {janela} GROUP BY 1
        """, (ini, fim))
        horas = dict(cur.fetchall())
        por_hora = [{"hora": h, "total": horas.get(h, 0)} for h in range(24)]

        def top(coluna, limite=8):
            cur.execute(f"""
                SELECT COALESCE(NULLIF(a.{coluna}, ''), '(sem)'), COUNT(*)
                FROM sd_atendimentos a WHERE {janela}
                GROUP BY 1 ORDER BY 2 DESC, 1 LIMIT %s
            """, (ini, fim, limite))
            return [{"nome": r[0], "total": r[1]} for r in cur.fetchall()]

        # Pausa: soma do tempo em pausa dentro do período (evento aberto conta até agora).
        cur.execute("""
            SELECT usuario_id,
                   SUM(EXTRACT(EPOCH FROM (LEAST(COALESCE(fim, NOW()::timestamp), %s) - GREATEST(inicio, %s))))
            FROM sd_status_eventos
            WHERE status = 'pausa' AND inicio < %s AND COALESCE(fim, NOW()::timestamp) > %s
            GROUP BY usuario_id
        """, (fim, ini, fim, ini))
        pausa_seg = {r[0]: float(r[1] or 0) for r in cur.fetchall()}

        cur.execute(f"""
            SELECT a.usuario_id, COUNT(*), COUNT(*) FILTER (WHERE {resolvido}),
                   COUNT(*) FILTER (WHERE t.finaliza),
                   COUNT(*) FILTER (WHERE COALESCE(a.fila_destino,'') <> ''),
                   COUNT(DISTINCT a.criado_em::date)
            FROM sd_atendimentos a LEFT JOIN sd_status_tipos t ON t.chave = a.status
            WHERE {janela} GROUP BY a.usuario_id
        """, (ini, fim))
        producao = {r[0]: r[1:] for r in cur.fetchall()}
        agora_status = {"online": 0, "pausa": 0, "indisponivel": 0, "offline": 0}
        por_operador = []
        for uid in _ids_equipe(cur):
            op = _operador(cur, uid)
            st = _status_atual(cur, uid)["status"]
            agora_status[st] = agora_status.get(st, 0) + 1
            n, res, fin, redir, dias_trab = producao.get(uid, (0, 0, 0, 0, 0))
            por_operador.append({
                "id": uid, "nome": op["nome"], "nivel": op["nivel"], "status": st,
                "total": n, "fcr_pct": round(100 * res / fin) if fin else None,
                "redirecionados": redir, "dias_com_atendimento": dias_trab,
                "media_por_dia": round(n / dias_trab, 1) if dias_trab else 0,
                "pausa_min": round(pausa_seg.get(uid, 0) / 60),
            })
        por_operador.sort(key=lambda o: (-o["total"], o["nome"]))
        categorias, filas, canais = top("categoria"), top("fila_entrada"), top("canal", 6)
        cur.execute(f"""
            SELECT a.status, COALESCE(t.rotulo, a.status), COUNT(*)
            FROM sd_atendimentos a LEFT JOIN sd_status_tipos t ON t.chave = a.status
            WHERE {janela} GROUP BY 1, 2, t.ordem ORDER BY t.ordem NULLS LAST, 3 DESC
        """, (ini, fim))
        por_status = [{"chave": r[0], "rotulo": r[1], "total": r[2]} for r in cur.fetchall()]
        cur.close()
        return {
            "periodo": {"chave": periodo, "inicio": d0.isoformat(), "fim": hoje.isoformat(), "dias": dias},
            "kpis": {
                "total": total, "media_dia": round(total / dias, 1),
                "fcr_pct": round(100 * resolvidos / finalizados) if finalizados else None,
                "redirecionados_pct": round(100 * redirecionados / total) if total else None,
                "em_aberto": em_aberto, "operadores_com_atendimento": operadores,
                "pausa_horas": round(sum(pausa_seg.values()) / 3600, 1),
                "agora": agora_status,
            },
            "por_dia": por_dia, "por_hora": por_hora, "por_operador": por_operador,
            "categorias": categorias, "filas": filas, "canais": canais, "status": por_status,
        }
    finally:
        conn.close()


# ── Kanban com carga (supervisor) ────────────────────────────────────────
# Mesma régua visual do "Carga da equipe" do painel de tarefas (main.py,
# CARGA_NIVEIS), mas a conta é do SD: pontos = atendimentos em aberto, e o
# que está parado há mais de 4h pesa 1,5x.
CARGA_SD_NIVEIS = {
    "tranquilo": {"rotulo": "Tranquilo",  "cor": "#0E9F6E", "icone": "ph-smiley"},
    "moderado":  {"rotulo": "Fluindo",    "cor": "#06D7E6", "icone": "ph-gauge"},
    "pesado":    {"rotulo": "Carga alta", "cor": "#F59E0B", "icone": "ph-warning"},
    "atolado":   {"rotulo": "Atolado",    "cor": "#EF4444", "icone": "ph-fire"},
}
CARGA_SD_LIMITES = {"moderado": 2, "pesado": 5, "atolado": 8}
PARADO_APOS = timedelta(hours=4)


def nivel_carga_sd(pontos: float) -> str:
    if pontos >= CARGA_SD_LIMITES["atolado"]:
        return "atolado"
    if pontos >= CARGA_SD_LIMITES["pesado"]:
        return "pesado"
    if pontos >= CARGA_SD_LIMITES["moderado"]:
        return "moderado"
    return "tranquilo"


@router.get("/api/sd/kanban")
def kanban(periodo: str = "hoje", faiston_token: str = Cookie(None)):
    """Cards do Kanban do supervisor: tudo que está em aberto (de qualquer
    data) + o que foi redirecionado/concluído no período. Junto vem a carga
    de cada operador pro painel "Carga da equipe"."""
    _supervisor(faiston_token)
    dias = PERIODOS_DASHBOARD.get(periodo)
    if not dias:
        raise HTTPException(status_code=400, detail="Período inválido")
    ref = agora()
    ini = datetime.combine(ref.date() - timedelta(days=dias - 1), time(0))
    conn = _conn()
    try:
        cur = conn.cursor()
        _encerrar_esquecidos(cur)
        conn.commit()
        # Em aberto vêm todos (são a base da carga da equipe); só os
        # finalizados do período têm teto. Com um LIMIT único, um mês cheio
        # de concluídos empurrava abertos antigos pra fora -- sumiam da
        # coluna e o operador sobrecarregado aparecia como "Tranquilo".
        base = f"""
            SELECT {_COLUNAS}, COALESCE(t.finaliza, FALSE)
            FROM sd_atendimentos a JOIN usuarios u ON u.id = a.usuario_id
            LEFT JOIN sd_status_tipos t ON t.chave = a.status
        """
        cur.execute(base + " WHERE NOT COALESCE(t.finaliza, FALSE) ORDER BY a.criado_em DESC")
        linhas = cur.fetchall()
        cur.execute(base + " WHERE COALESCE(t.finaliza, FALSE) AND a.criado_em >= %s ORDER BY a.criado_em DESC LIMIT 600",
                    (ini,))
        linhas += cur.fetchall()
        itens = []
        for r in linhas:
            item = _linha(r[:-1])
            item["finaliza"] = r[-1]
            item["coluna"] = ("aberto" if not r[-1] else
                              "redirecionado" if item["fila_destino"] else "concluido")
            item["idade_min"] = max(0, int((ref - r[13]).total_seconds() // 60))
            item["parado"] = not r[-1] and (ref - r[13]) > PARADO_APOS
            itens.append(item)
        carga = []
        for uid in _ids_equipe(cur):
            op = _operador(cur, uid)
            meus = [i for i in itens if i["usuario_id"] == uid and i["coluna"] == "aberto"]
            parados = sum(1 for i in meus if i["parado"])
            pontos = round(len(meus) + 0.5 * parados, 1)
            nivel = nivel_carga_sd(pontos)
            j_ini, j_fim, _ = _jornada_em(cur, op, ref)
            c_ini, c_fim = janela_contagem(ref, j_ini, j_fim)
            cur.execute("SELECT COUNT(*) FROM sd_atendimentos WHERE usuario_id = %s AND criado_em >= %s AND criado_em < %s",
                        (uid, c_ini, c_fim))
            turno_total = cur.fetchone()[0]
            carga.append({
                "usuario_id": uid, "nome": op["nome"], "nivel_op": op["nivel"],
                "status": _status_atual(cur, uid)["status"],
                "abertos": len(meus), "parados": parados, "pontos": pontos,
                "nivel": nivel, **{k: CARGA_SD_NIVEIS[nivel][k] for k in ("rotulo", "cor", "icone")},
                "limite_atolado": CARGA_SD_LIMITES["atolado"],
                "turno_total": turno_total, "meta": op["meta_turno"],
            })
        carga.sort(key=lambda p: (-p["pontos"], p["nome"]))
        cur.close()
        return {
            "periodo": {"chave": periodo, "inicio": ini.date().isoformat(), "dias": dias},
            "itens": itens, "carga": carga, "niveis": CARGA_SD_NIVEIS, "limites": CARGA_SD_LIMITES,
            "resumo": {
                "atolados": sum(1 for p in carga if p["nivel"] == "atolado"),
                "pesados": sum(1 for p in carga if p["nivel"] == "pesado"),
                "livres": sum(1 for p in carga if not p["abertos"]),
                "abertos": sum(p["abertos"] for p in carga),
                "parados": sum(p["parados"] for p in carga),
            },
        }
    finally:
        conn.close()


class MoverAtendimentoModel(BaseModel):
    status: str
    fila_destino: Optional[str] = Field(default=None, max_length=120)


@router.patch("/api/sd/atendimentos/{aid}/status")
def mudar_status(aid: int, body: MoverAtendimentoModel, faiston_token: str = Cookie(None)):
    """Troca o status (arrastar card no Kanban do supervisor).

    A coluna Redirecionado/Concluído e o FCR vêm de `fila_destino`, não do
    status -- então mexer só no status fazia o card voltar pra coluna antiga
    no reload. Concluir limpa a fila de destino (resolvido aqui); redirecionar
    exige saber pra qual fila."""
    sess = _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        atual = _buscar(cur, aid)
        if not _pode_editar(sess, atual):
            raise HTTPException(status_code=403, detail="Só dá pra alterar os próprios atendimentos")
        cur.execute("SELECT 1 FROM sd_status_tipos WHERE chave = %s AND ativo", (body.status,))
        if not cur.fetchone():
            raise HTTPException(status_code=400, detail="Status inválido")
        fila = atual["fila_destino"] or ""
        if body.status == "concluido":
            fila = ""
        elif body.status == "redirecionado":
            fila = (body.fila_destino if body.fila_destino is not None else fila).strip()
            if not fila or fila == atual["fila_entrada"]:
                raise HTTPException(status_code=400, detail="Informe a fila para onde o atendimento foi redirecionado")
            _garantir_opcao(cur, "sd_filas", fila)
        cur.execute("UPDATE sd_atendimentos SET status = %s, fila_destino = %s, atualizado_em = NOW() WHERE id = %s",
                    (body.status, fila, aid))
        conn.commit()
        out = _buscar(cur, aid)
        cur.close()
        return out
    finally:
        conn.close()


# ── Equipe (supervisor) ──────────────────────────────────────────────────
def _ids_equipe(cur) -> List[int]:
    cur.execute("""
        SELECT id FROM usuarios
        WHERE ativo AND (time = %s OR cargo IN %s)
        ORDER BY nome
    """, (TIME_SERVICE_DESK, CARGOS_SD))
    return [r[0] for r in cur.fetchall()]


@router.get("/api/sd/equipe")
def equipe(faiston_token: str = Cookie(None)):
    """Visão do supervisor: status em tempo real e produção do turno de
    cada operador."""
    _supervisor(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        _encerrar_esquecidos(cur)
        conn.commit()
        ref = agora()
        out = []
        for uid in _ids_equipe(cur):
            op = _operador(cur, uid)
            j_ini, j_fim, _ = _jornada_em(cur, op, ref)
            ini, fim = janela_contagem(ref, j_ini, j_fim)
            out.append({
                **{k: v for k, v in op.items() if k not in ("jornada_inicio", "jornada_fim")},
                "jornada_inicio": _hhmm(op["jornada_inicio"]), "jornada_fim": _hhmm(op["jornada_fim"]),
                "status": _status_atual(cur, uid),
                "kpis": _kpis(cur, uid, ini, fim, op["meta_turno"]),
            })
        cur.close()
        return out
    finally:
        conn.close()


class OperadorModel(BaseModel):
    matricula: str = Field(default="", max_length=30)
    nivel: str = "N1"
    site: str = Field(default="", max_length=60)
    jornada_inicio: str = ""
    jornada_fim: str = ""
    regime: str = Field(default="", max_length=20)
    meta_turno: int = Field(default=20, ge=0, le=500)


@router.put("/api/sd/operadores/{uid}")
def salvar_operador(uid: int, body: OperadorModel, faiston_token: str = Cookie(None)):
    _supervisor(faiston_token)
    if body.nivel not in ("N1", "N2", "N3"):
        raise HTTPException(status_code=400, detail="Nível inválido")
    conn = _conn()
    try:
        cur = conn.cursor()
        _operador(cur, uid)  # 404 se não existe
        cur.execute("""
            INSERT INTO sd_operadores (usuario_id, matricula, nivel, site, jornada_inicio, jornada_fim, regime, meta_turno)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
            ON CONFLICT (usuario_id) DO UPDATE SET matricula=EXCLUDED.matricula, nivel=EXCLUDED.nivel,
                site=EXCLUDED.site, jornada_inicio=EXCLUDED.jornada_inicio, jornada_fim=EXCLUDED.jornada_fim,
                regime=EXCLUDED.regime, meta_turno=EXCLUDED.meta_turno, atualizado_em=NOW()
        """, (uid, body.matricula.strip(), body.nivel, body.site.strip(), _hora(body.jornada_inicio),
              _hora(body.jornada_fim), body.regime.strip(), body.meta_turno))
        conn.commit(); cur.close()
        return {"sucesso": True}
    finally:
        conn.close()


# ── Escalas ──────────────────────────────────────────────────────────────
@router.get("/api/sd/escalas")
def listar_escalas(inicio: str = "", dias: int = Query(7, ge=1, le=62), faiston_token: str = Cookie(None)):
    _sessao(faiston_token)
    d0 = _data(inicio, "inicio") or (agora().date() - timedelta(days=agora().weekday()))
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT e.id, e.data, e.usuario_id, COALESCE(u.nome, ''), e.rotulo, e.hora_inicio, e.hora_fim, e.tipo, e.obs
            FROM sd_escalas e LEFT JOIN usuarios u ON u.id = e.usuario_id
            WHERE e.data >= %s AND e.data < %s ORDER BY e.data, e.hora_inicio, u.nome
        """, (d0, d0 + timedelta(days=dias)))
        itens = [{"id": r[0], "data": r[1].isoformat(), "usuario_id": r[2], "nome": r[3] or r[4],
                  "hora_inicio": _hhmm(r[5]), "hora_fim": _hhmm(r[6]), "tipo": r[7], "obs": r[8] or "",
                  "noturno": r[5] >= r[6]} for r in cur.fetchall()]
        cur.close()
        return {"inicio": d0.isoformat(), "dias": dias, "itens": itens}
    finally:
        conn.close()


class EscalaModel(BaseModel):
    data: str
    usuario_id: Optional[int] = None
    rotulo: str = Field(default="", max_length=80)
    hora_inicio: str
    hora_fim: str
    tipo: str = "turno"
    obs: str = Field(default="", max_length=200)


@router.post("/api/sd/escalas")
def criar_escala(body: EscalaModel, faiston_token: str = Cookie(None)):
    _supervisor(faiston_token)
    if body.tipo not in TIPOS_ESCALA:
        raise HTTPException(status_code=400, detail="Tipo de escala inválido")
    if not body.usuario_id and not body.rotulo.strip():
        raise HTTPException(status_code=400, detail="Informe o operador ou um rótulo (ex.: RJ Plantão)")
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO sd_escalas (data, usuario_id, rotulo, hora_inicio, hora_fim, tipo, obs)
            VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id
        """, (_data(body.data), body.usuario_id, body.rotulo.strip(), _hora(body.hora_inicio),
              _hora(body.hora_fim), body.tipo, body.obs.strip()))
        eid = cur.fetchone()[0]
        conn.commit(); cur.close()
        return {"sucesso": True, "id": eid}
    finally:
        conn.close()


class Gerar12x36Model(BaseModel):
    usuario_id: int
    data_inicio: str
    data_fim: str
    hora_inicio: str = "19:00"
    hora_fim: str = "07:00"
    substituir: bool = True


@router.post("/api/sd/escalas/gerar-12x36")
def gerar_12x36(body: Gerar12x36Model, faiston_token: str = Cookie(None)):
    """Preenche um período inteiro de 12x36 (dia sim, dia não) de uma vez --
    em vez de lançar plantão por plantão."""
    _supervisor(faiston_token)
    d0, d1 = _data(body.data_inicio, "data_inicio"), _data(body.data_fim, "data_fim")
    if d1 < d0 or (d1 - d0).days > 92:
        raise HTTPException(status_code=400, detail="Período inválido (máximo 3 meses)")
    h0, h1 = _hora(body.hora_inicio), _hora(body.hora_fim)
    conn = _conn()
    try:
        cur = conn.cursor()
        if body.substituir:
            cur.execute("DELETE FROM sd_escalas WHERE usuario_id = %s AND data BETWEEN %s AND %s AND tipo = 'turno'",
                        (body.usuario_id, d0, d1))
        n = 0
        for dia in dias_12x36(d0, d1):
            cur.execute("""
                INSERT INTO sd_escalas (data, usuario_id, hora_inicio, hora_fim, tipo) VALUES (%s,%s,%s,%s,'turno')
            """, (dia, body.usuario_id, h0, h1))
            n += 1
        conn.commit(); cur.close()
        return {"sucesso": True, "criados": n}
    finally:
        conn.close()


@router.delete("/api/sd/escalas/{eid}")
def excluir_escala(eid: int, faiston_token: str = Cookie(None)):
    _supervisor(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM sd_escalas WHERE id = %s", (eid,))
        conn.commit(); cur.close()
        return {"sucesso": True}
    finally:
        conn.close()


@router.get("/api/sd/operadores")
def listar_operadores(faiston_token: str = Cookie(None)):
    """Lista enxuta pros selects de escala (qualquer um da equipe vê)."""
    _sessao(faiston_token)
    conn = _conn()
    try:
        cur = conn.cursor()
        out = []
        for uid in _ids_equipe(cur):
            op = _operador(cur, uid)
            out.append({"id": uid, "nome": op["nome"], "nivel": op["nivel"], "matricula": op["matricula"],
                        "jornada_inicio": _hhmm(op["jornada_inicio"]), "jornada_fim": _hhmm(op["jornada_fim"])})
        cur.close()
        return out
    finally:
        conn.close()
